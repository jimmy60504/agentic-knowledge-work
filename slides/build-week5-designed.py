#!/usr/bin/env python3
"""第五週投影片排版版：把 drafts/15-week5-slides.md 產成排過版的 PPTX（slides/week5-designed.pptx）。

用法：python3 slides/build-week5-designed.py

與素版 build-week5.py 共用逐頁稿解析，差別在版面：
- 自繪示意圖放在右側視覺區，保持原比例、不加框。
- 來源截圖做成斜放的白色卡片（細框加陰影），截圖統一裁成 16:10，保留上方的標題區。
  一到兩張：各自斜放，下方放案例的簡短說明與來源。
  三張以上：卡片錯開堆疊，下方改為編號清單，逐張列出說明與來源。
- 示意圖與截圖並存時，示意圖在左、卡片在右。
- 只有文字的頁面：主訊息放在象牙色色帶，要點排成兩欄卡片。
- 文字放不下時自動縮小字級，並在終端機提示。
"""
import importlib.util
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
spec = importlib.util.spec_from_file_location("bw4", HERE / "build-week4.py")
bw4 = importlib.util.module_from_spec(spec)
sys.path.insert(0, str(HERE))
spec.loader.exec_module(bw4)

SRC = ROOT / "drafts/15-week5-slides.md"
OUT = HERE / "week5-designed.pptx"

DARK = RGBColor(0x2A, 0x26, 0x23)
MUTED = RGBColor(0x6B, 0x65, 0x60)
ACCENT = RGBColor(0xBE, 0x89, 0x79)
BLUE = RGBColor(0xA6, 0xB7, 0xBA)
IVORY = RGBColor(0xF5, 0xF2, 0xEC)
GREY = RGBColor(0xD8, 0xD2, 0xC9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "PingFang TC"
W, H, M = 13.333, 7.5, 0.7
CARD_ASPECT = 1.6
ANGLES = [-3.5, 2.5, -2.0, 3.0, -1.5]
# 示意圖已包含要點文字的頁面：畫面只留主訊息，要點移到講者備註，示意圖放大
DIAGRAM_CARRIES = set()  # 使用者 2026-09-29 要求示意圖頁面保留內文，目前不省略任何頁的要點

notes = bw4.notes


# ---------- 基本元件 ----------

def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    shp.shadow.inherit = False
    return shp


def shadow(shp, alpha=24000, blur=90000, dist=30000):
    spPr = shp._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(blur), dist=str(dist),
                          dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val="000000")
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha))


def runs(slide, x, y, w, h, paras, align="l", anchor="t"):
    """paras：[(文字, 字級, 粗體, 顏色, 段後點數)]。"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    for i, (text, size, bold, color, after) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        p.space_after = Pt(after)
        p.line_spacing = 1.15
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        f.name = FONT
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", FONT)
    return box


def est_h(text, size, w, spacing=1.2):
    per_line = max(int(w * 72 / (size * 1.02)), 1)
    lines = -(-len(text) // per_line)
    return lines * size / 72 * spacing * 1.15


# ---------- 解析頁面 ----------

def images_of(page):
    out = []
    for line in page["notes"]:
        if line.startswith("圖："):
            body = line[len("圖："):].strip()
            rel, _, cap = body.partition("｜")
            rel = rel.strip().strip("`")
            if rel and rel not in [r for r, _ in out] and (ROOT / rel).exists():
                out.append((rel, cap.strip()))
    page["notes"] = [l for l in page["notes"] if not l.startswith("圖：")]
    return out


def is_source(rel):
    return "week5-sources" in rel


def split_cap(cap):
    """「來源，日期：說明」拆成 (說明, 來源)。"""
    if "：" in cap:
        src, desc = cap.split("：", 1)
        return desc.strip(), src.strip()
    return cap.strip(), ""


def parse_bullets(page):
    lead = [v for k, v in page["blocks"] if k == "lead"]
    items = []
    for k, v in page["blocks"]:
        if k not in ("bullet", "num"):
            continue
        m = re.match(r"^\*\*(.+?)\*\*[：:]\s*(.*)$", v)
        items.append((m.group(1), m.group(2)) if m else ("", v))
    return (lead[0] if lead else ""), items


# ---------- 視覺元件 ----------

def picture(slide, rel, x, y, w, h, crop_to=None):
    path = ROOT / rel
    iw, ih = Image.open(path).size
    crop_b = 0.0
    if crop_to and iw / ih < crop_to:
        crop_b = 1 - (iw / crop_to) / ih
        ih = iw / crop_to
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px, py = x + (w - pw) / 2, y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    if crop_b:
        pic.crop_bottom = crop_b
    return pic, (px, py, pw, ph)


def card(slide, rel, cx, cy, cw, angle):
    """以中心 (cx, cy)、寬 cw 放一張斜放的截圖卡片，回傳卡片外框 (x, y, w, h)。"""
    pad = 0.07
    ch = (cw - 2 * pad) / CARD_ASPECT + 2 * pad
    x, y = cx - cw / 2, cy - ch / 2
    bg = rect(slide, x, y, cw, ch, WHITE, line=GREY)
    shadow(bg)
    bg.rotation = angle
    pic, _ = picture(slide, rel, x + pad, y + pad, cw - 2 * pad, ch - 2 * pad, crop_to=CARD_ASPECT)
    pic.rotation = angle
    return x, y, cw, ch


def cards_single(slide, srcs, x, y, w, h):
    """一到兩張：各自斜放，下方放說明與來源。"""
    n = len(srcs)
    gap = 0.35
    colw = (w - gap * (n - 1)) / n
    cap_h = 0.8
    for i, (rel, cap) in enumerate(srcs):
        cx0 = x + i * (colw + gap)
        cw = min(colw * 0.88, (h - cap_h - 0.3) * CARD_ASPECT)
        ch = cw / CARD_ASPECT
        cy = y + (h - cap_h) / 2
        _, top, _, chh = card(slide, rel, cx0 + colw / 2, cy, cw, ANGLES[i])
        desc, src = split_cap(cap)
        paras = [(desc, 12.5, True, DARK, 2)]
        if src:
            paras.append((src, 10, False, MUTED, 0))
        runs(slide, cx0 + (colw - cw) / 2, top + chh + 0.2, cw, cap_h, paras, align="l")


def cards_stack(slide, srcs, x, y, w, h):
    """三張以上：錯開堆疊，下方編號清單。"""
    n = len(srcs)
    item_h = [0.2 + est_h(split_cap(c)[0] + "（" + split_cap(c)[1] + "）", 10.5, w - 0.4) for _, c in srcs]
    list_h = min(sum(item_h) + 0.1, h * 0.46)
    area_h = h - list_h - 0.15
    cw = min(w * 0.66, (area_h * 0.72) * CARD_ASPECT)
    ch = cw / CARD_ASPECT
    span_x = w - cw - 0.2
    span_y = area_h - ch - 0.15
    for i, (rel, cap) in enumerate(srcs):
        t = i / (n - 1)
        cx = x + 0.1 + cw / 2 + span_x * t
        cy = y + 0.1 + ch / 2 + span_y * t
        cx0, cy0, _, _ = card(slide, rel, cx, cy, cw, ANGLES[i % len(ANGLES)])
        # 卡片左上角的編號
        dot = rect(slide, cx0 - 0.12, cy0 - 0.12, 0.34, 0.34, ACCENT, shape=MSO_SHAPE.OVAL)
        tf = dot.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(i + 1)
        r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    paras = []
    for i, (_, cap) in enumerate(srcs):
        desc, src = split_cap(cap)
        paras.append((f"{i + 1}  {desc}" + (f"（{src}）" if src else ""), 10.5, False, DARK, 3))
    runs(slide, x, y + h - list_h, w, list_h, paras)


def diagrams(slide, diags, x, y, w, h):
    if len(diags) == 1:
        picture(slide, diags[0][0], x, y, w, h)
        return
    aspects = [Image.open(ROOT / r).size[0] / Image.open(ROOT / r).size[1] for r, _ in diags]
    gap = 0.3
    if all(a > 1.4 for a in aspects):
        hh = (h - gap) / 2
        for i, (r, _) in enumerate(diags[:2]):
            picture(slide, r, x, y + i * (hh + gap), w, hh)
    else:
        ww = (w - gap) / 2
        for i, (r, _) in enumerate(diags[:2]):
            picture(slide, r, x + i * (ww + gap), y, ww, h)


def visuals(slide, imgs, x, y, w, h):
    diags = [i for i in imgs if not is_source(i[0])]
    srcs = [i for i in imgs if is_source(i[0])]
    if diags and srcs:
        dw = w * (0.46 if len(srcs) <= 2 else 0.36)
        diagrams(slide, diags, x, y, dw, h)
        x, w = x + dw + 0.3, w - dw - 0.3
    elif diags:
        diagrams(slide, diags, x, y, w, h)
        return
    if len(srcs) <= 2:
        cards_single(slide, srcs, x, y, w, h)
    else:
        cards_stack(slide, srcs, x, y, w, h)


# ---------- 頁面 ----------

def header(slide, section, title):
    if section:
        rect(slide, M, 0.43, 0.1, 0.1, ACCENT)
        runs(slide, M + 0.2, 0.3, 8, 0.35, [(section, 11, False, MUTED, 0)])
    runs(slide, M, 0.62, W - 2 * M, 0.7, [(title, 30, True, DARK, 0)])
    rect(slide, M, 1.36, 0.8, 0.05, ACCENT)


def footer(slide, page, total, sources):
    if sources:
        runs(slide, M, H - 0.5, W - 2 * M - 1.0, 0.4, [("來源：" + "；".join(sources), 9, False, MUTED, 0)])
    runs(slide, W - M - 0.8, H - 0.5, 0.8, 0.3, [(f"{page['n']} / {total}", 10, False, MUTED, 0)], align="r")


def text_column(slide, lead, items, x, y, w, bottom, page):
    for scale in (1.0, 0.92, 0.85, 0.78):
        ls, ks, bs = 20 * scale, 16.5 * scale, 13 * scale
        need = est_h(lead, ls, w) + 0.35 if lead else 0
        for k, b in items:
            need += (est_h(k, ks, w) if k else 0) + (est_h(b, bs, w) if b else 0) + 0.24
        if y + need <= bottom:
            break
    else:
        print(f"  [文字偏多] 第 {page['n']} 頁「{page['title']}」")
    if not items:
        ls = 24
    if lead:
        lh = est_h(lead, ls, w)
        runs(slide, x, y, w, lh + 0.1, [(lead, ls, True, DARK, 0)])
        y += lh + 0.35
    for k, b in items:
        kh = est_h(k, ks, w) if k else 0
        bh = est_h(b, bs, w) if b else 0
        rect(slide, x, y + 0.05, 0.05, kh + bh + 0.02, ACCENT if k else GREY)
        paras = []
        if k:
            paras.append((k, ks, True, DARK, 1))
        if b:
            paras.append((b, bs, False, MUTED, 0))
        runs(slide, x + 0.18, y, w - 0.18, kh + bh + 0.1, paras)
        y += kh + bh + 0.24


def text_only(slide, lead, items, page):
    y = 1.75
    if lead:
        lh = est_h(lead, 24, W - 2 * M - 0.6)
        band = rect(slide, M, y, W - 2 * M, lh + 0.5, IVORY)
        rect(slide, M, y, 0.08, lh + 0.5, ACCENT)
        runs(slide, M + 0.35, y + 0.22, W - 2 * M - 0.6, lh + 0.1, [(lead, 24, True, DARK, 0)])
        y += lh + 0.85
    if not items:
        return
    if any("／" in body for _, body in items):
        question_panels(slide, items, y)
        return
    quote = [b for k, b in items if not k and b.startswith("課堂提問")]
    items = [(k, b) for k, b in items if not (not k and b.startswith("課堂提問"))]
    if quote:
        qt = quote[0].split("：", 1)[-1]
        qy = H - 0.75 - 1.75
        rect(slide, M, qy, W - 2 * M, 1.55, WHITE, line=ACCENT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        runs(slide, M + 0.4, qy + 0.12, 3, 0.4, [("課堂提問", 13, True, ACCENT, 0)])
        runs(slide, M + 0.4, qy + 0.45, W - 2 * M - 0.8, 1.0, [(qt, 24, True, DARK, 0)], anchor="m")
        bottom_limit = qy - 0.3
    else:
        bottom_limit = H - 0.75
    if not items:
        return
    cols = 2 if len(items) >= 3 or any(k for k, _ in items) else 1
    gap = 0.3
    cw = (W - 2 * M - gap * (cols - 1)) / cols
    rows = -(-len(items) // cols)
    avail = bottom_limit - y
    rh = min(2.2 if len(items) <= 2 else 1.5, (avail - gap * (rows - 1)) / rows)
    big = len(items) <= 2
    for i, (k, b) in enumerate(items):
        r, c = divmod(i, cols)
        cx, cy = M + c * (cw + gap), y + r * (rh + gap)
        box = rect(slide, cx, cy, cw, rh, WHITE, line=GREY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        rect(slide, cx, cy + 0.2, 0.07, rh - 0.4, ACCENT)
        paras = []
        if k:
            paras.append((k, 22 if big else 17, True, DARK, 4))
        if b:
            paras.append((b, 16 if big else 13.5, False, MUTED, 0))
        runs(slide, cx + 0.3, cy + 0.1, cw - 0.5, rh - 0.2, paras, anchor="m")


def question_panels(slide, items, y):
    """題目以「／」分隔時：每組一欄，題目逐行加編號。"""
    gap = 0.35
    cols = len(items)
    cw = (W - 2 * M - gap * (cols - 1)) / cols
    ph = H - 0.75 - y
    for c, (k, body) in enumerate(items):
        cx = M + c * (cw + gap)
        rect(slide, cx, y, cw, ph, WHITE, line=GREY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
        rect(slide, cx, y, cw, 0.62, IVORY, shape=MSO_SHAPE.RECTANGLE)
        rect(slide, cx, y, 0.08, 0.62, ACCENT)
        runs(slide, cx + 0.3, y + 0.1, cw - 0.5, 0.45, [(k, 18, True, DARK, 0)], anchor="m")
        qs = [q.strip() for q in body.split("／") if q.strip()]
        size = 15 if len(qs) <= 6 else 14
        paras = [(f"{i + 1}.  {q}？" if not q.endswith("？") else f"{i + 1}.  {q}", size, False, DARK, 7)
                 for i, q in enumerate(qs)]
        runs(slide, cx + 0.3, y + 0.8, cw - 0.5, ph - 0.9, paras)


def intermission(prs, page, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, IVORY)
    cover_img = ROOT / "assets/week5-visuals/w5-01-abstract.png"
    if cover_img.exists():
        pic = s.shapes.add_picture(str(cover_img), Inches(W * 0.55), 0, Inches(W * 0.45), Inches(H))
        iw, ih = Image.open(cover_img).size
        want = (W * 0.45) / H
        have = iw / ih
        if have > want:
            cut = (1 - want / have)
            pic.crop_left = cut / 2
            pic.crop_right = cut / 2
    lead, items = parse_bullets(page)
    runs(s, M, 0.9, 6.5, 1.0, [("中場休息", 44, True, DARK, 0)])
    rect(s, M + 0.05, 1.95, 1.0, 0.06, ACCENT)
    runs(s, M, 2.2, 6.8, 0.6, [(lead, 20, True, DARK, 0)])
    y = 3.0
    for k, b in items:
        if not k:
            runs(s, M, H - 1.0, 7, 0.5, [(b, 13, False, MUTED, 0)])
            continue
        rect(s, M, y, 6.6, 0.95, WHITE, line=GREY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        rect(s, M, y + 0.18, 0.07, 0.6, ACCENT)
        runs(s, M + 0.3, y + 0.08, 6.1, 0.8, [(k + "？", 19, True, DARK, 2), (b, 12.5, False, MUTED, 0)], anchor="m")
        y += 1.1
    notes(s, page["notes"])


def content(prs, page, total, section):
    if page["title"] == "中場休息":
        intermission(prs, page, total)
        return
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, WHITE)
    imgs = images_of(page)
    lead, items = parse_bullets(page)
    sources = [l[len("來源："):].strip() for l in page["notes"] if l.startswith("來源：")]
    header(s, section, page["title"])
    col_w = 5.2
    if imgs and page["title"] in DIAGRAM_CARRIES:
        page["notes"] = ["【畫面要點（已由示意圖呈現）】"] + [f"{k}：{b}" if k else b for k, b in items] + page["notes"]
        items = []
        col_w = 3.9
    wide = (len(imgs) == 1 and not is_source(imgs[0][0])
            and (lambda sz: sz[0] / sz[1])(Image.open(ROOT / imgs[0][0]).size) > 1.8)
    if wide:
        # 寬圖：文字在上方兩欄，圖橫跨下方
        if lead:
            runs(s, M, 1.7, W - 2 * M, 0.6, [(lead, 20, True, DARK, 0)])
        half = (W - 2 * M - 0.4) / 2
        for i, (k, bd) in enumerate(items[:2]):
            x = M + i * (half + 0.4)
            rect(s, x, 2.45, 0.05, 0.62, ACCENT)
            runs(s, x + 0.18, 2.4, half - 0.2, 0.8, [(k, 16.5, True, DARK, 1), (bd, 13, False, MUTED, 0)])
        picture(s, imgs[0][0], M, 3.45, W - 2 * M, H - 0.8 - 3.45)
    elif imgs:
        text_column(s, lead, items, M, 1.75, col_w, H - 0.75, page)
        vx = M + col_w + 0.45
        visuals(s, imgs, vx, 1.7, W - M - vx, H - 0.75 - 1.7)
    else:
        text_only(s, lead, items, page)
    footer(s, page, total, sources)
    notes(s, page["notes"])


def cover(prs, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    imgs = images_of(page)
    rect(s, 0, 0, W, H, IVORY)
    if imgs:
        s.shapes.add_picture(str(ROOT / imgs[0][0]), 0, 0, Inches(W), Inches(H))
    lines = [v for k, v in page["blocks"] if k in ("bullet", "lead")]
    title = lines[0] if lines else page["title"]
    runs(s, M + 0.2, 2.6, 9.0, 1.2, [(title, 40, True, DARK, 0)])
    if len(lines) > 1:
        runs(s, M + 0.2, 4.35, 6.6, 0.8, [(lines[1], 18, False, MUTED, 0)])
    rect(s, M + 0.25, 4.2, 1.0, 0.06, ACCENT)
    notes(s, page["notes"])


def divider(prs, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, IVORY)
    title = page["title"]
    num, _, name = title.partition("、")
    digits = {"一": "01", "二": "02", "三": "03", "四": "04", "五": "05", "六": "06", "七": "07", "八": "08"}
    if name:
        runs(s, M, 1.9, 5, 2.2, [(digits.get(num, num), 120, True, ACCENT, 0)])
        rect(s, M + 0.05, 4.15, 1.2, 0.06, ACCENT)
        runs(s, M, 4.35, W - 2 * M, 1.0, [(name, 36, True, DARK, 0)])
    else:
        rect(s, M, 3.0, 1.2, 0.06, ACCENT)
        runs(s, M, 3.2, W - 2 * M, 1.2, [(title, 36, True, DARK, 0)])
    notes(s, page["notes"])


def main():
    pages = bw4.parse(SRC)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    total = sum(1 for p in pages if not p.get("divider"))
    section = ""
    for p in pages:
        if p.get("divider"):
            divider(prs, p)
            section = p["title"]
            continue
        if p["n"] == 1:
            cover(prs, p)
        else:
            content(prs, p, total, section)
    prs.save(OUT)
    print(f"{OUT.name}: {len(pages)} 頁（含 {len(pages) - total} 頁段落標題）← {SRC.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
