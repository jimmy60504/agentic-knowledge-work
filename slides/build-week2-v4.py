#!/usr/bin/env python3
"""第二週講師版投影片，依 drafts/08-week2-slide-copy.md 第七版建置。

用法：python3 slides/build-week2-v4.py  → slides/week2-agent-knowledge-work-v4.pptx
原則：標題是一句主張、畫面是證據、解釋放備註（kb/slide-design-principles-argument-to-slides.md）。
版面依 drafts/07-visual-design-rules.md；配色暖炭灰、陶土紅、鼠尾草綠；圖示 assets/week2-icons（Tabler，MIT）。
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
A = ROOT / "assets"
IC = A / "week2-icons"
OUT = ROOT / "slides" / "week2-agent-knowledge-work-v4.pptx"

DARK = RGBColor(0x2A, 0x26, 0x23)
DARK2 = RGBColor(0x3A, 0x35, 0x31)
ACCENT = RGBColor(0xB8, 0x50, 0x42)
SAGE = RGBColor(0xA7, 0xBE, 0xAE)
TINT = RGBColor(0xEE, 0xF2, 0xEF)
MUTED = RGBColor(0x6B, 0x65, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD4, 0xCE)
FONT = "PingFang TC"
W, H, M = 13.333, 7.5, 0.6

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


# ---------- 元件 ----------
def _font(run, size, bold=False, color=DARK):
    f = run.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def tb(s, x, y, w, h, lines, size=15, bold=False, color=DARK, align="l", anchor="t", spacing=5):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        d = ln if isinstance(ln, dict) else {"text": ln}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[d.get("align", align)]
        p.space_after = Pt(d.get("after", spacing))
        if d.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.2))))
            pPr.set("indent", str(-int(Inches(0.2))))
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        r = p.add_run()
        r.text = d["text"]
        _font(r, d.get("size", size), d.get("bold", bold), d.get("color", color))
    return box


def rect(s, x, y, w, h, fill=TINT, line=None, radius=0.08):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def oval(s, x, y, d, fill):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    return c


def num(s, x, y, n, d=0.42, fill=ACCENT, color=WHITE, size=14):
    c = oval(s, x, y, d, fill)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    _font(r, size, True, color)
    return c


def icon(s, name, x, y, d=0.6, variant="accent", circle=None):
    """單色圖示；circle 給底色時，圖示置於圓內。"""
    if circle is not None:
        oval(s, x, y, d, circle)
        pad = d * 0.22
        s.shapes.add_picture(str(IC / f"{name}-{variant}.png"), Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad))
    else:
        s.shapes.add_picture(str(IC / f"{name}-{variant}.png"), Inches(x), Inches(y), Inches(d), Inches(d))


def picture(s, path, x, y, w, h, align="l"):
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    pw, ph = (h * ar, h) if w / h > ar else (w, w / ar)
    px = x + (w - pw) / 2 if align == "c" else x
    py = y + (h - ph) / 2 if align == "c" else y
    pic = s.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    return px, py, pw, ph


def arrow(s, x1, y1, x2, y2, color=MUTED, dashed=False, width=1.5, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle"); t.set("w", "med"); t.set("h", "med")
    return c


def table(s, x, y, w, rows, col_w, size=14, row_h=0.5, first_col_bold=True):
    gt = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = gt.table
    tblPr = t._tbl.tblPr
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            head = ri == 0
            cell.fill.fore_color.rgb = DARK if head else (WHITE if ri % 2 else TINT)
            p = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            r = p.add_run(); r.text = str(val)
            _font(r, size, head or (first_col_bold and ci == 0), WHITE if head else DARK)
    return gt


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def assert_slide(title, seg, right=None):
    """白底主張頁：左上段落標，標題句 30pt 兩行內，右上可放時間。"""
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    tb(s, M, 0.45, 6, 0.3, seg, size=11, color=MUTED)
    if right:
        tb(s, W - M - 4, 0.45, 4, 0.3, right, size=11, color=MUTED, align="r")
    tb(s, M, 0.85, W - 2 * M, 1.2, title, size=30, bold=True)
    return s


def card(s, x, y, w, h, head, body, ic=None, head_size=17, body_size=14, n=None):
    rect(s, x, y, w, h)
    hx = x + 0.25
    if ic:
        icon(s, ic, x + 0.25, y + 0.25, 0.6, circle=WHITE)
        hx = x + 1.0
    elif n is not None:
        num(s, x + 0.25, y + 0.3, n)
        hx = x + 0.8
    tb(s, hx, y + 0.3, w - (hx - x) - 0.2, 0.5, head, size=head_size, bold=True, anchor="m")
    tb(s, x + 0.25, y + 1.0, w - 0.5, h - 1.15, body, size=body_size)


BODY_Y = 2.25  # 主張頁內容起始

# ============ 01 標題 ============
s = prs.slides.add_slide(BLANK); bg(s, DARK)
tb(s, M, 2.3, 11, 1.2, "先體驗 agent", size=48, bold=True, color=WHITE)
tb(s, M, 3.6, 11, 0.5, "Agent 時代的知識工作｜第二週", size=18, color=SAGE)
notes(s, "三段：前半段 30 分講 agent 能協助什麼；示範 30 分講師走一遍；練習 120 分自己做一份 10 到 12 頁的簡報。")

# ============ 02 研究專案 ============
s = assert_slide("研究專案裡的材料、想法與成果持續累積，agent 可以參與其中", "前半段")
items = [("book", "材料", "論文、資料、試過的方法"), ("bulb", "想法", "邊讀邊分析，隨時調整"), ("device-projector", "成果", "整理資料、小工具、向同事說明")]
cw, gap = 3.6, 0.55
x0 = (W - (3 * cw + 2 * gap)) / 2
for i, (ic, h1, body) in enumerate(items):
    x = x0 + i * (cw + gap)
    rect(s, x, 2.5, cw, 2.3)
    icon(s, ic, x + cw / 2 - 0.45, 2.75, 0.9, circle=WHITE)
    tb(s, x, 3.75, cw, 0.4, h1, size=20, bold=True, align="c")
    tb(s, x + 0.3, 4.2, cw - 0.6, 0.5, body, size=13, color=MUTED, align="c")
    if i < 2:
        arrow(s, x + cw + 0.08, 3.65, x + cw + gap - 0.08, 3.65, color=ACCENT, width=2)
cx3 = x0 + 2 * (cw + gap) + cw / 2
cx1 = x0 + cw / 2
arrow(s, cx3, 4.85, cx3, 5.45, color=MUTED, head=False)
arrow(s, cx3, 5.45, cx1, 5.45, color=MUTED, head=False)
arrow(s, cx1, 5.45, cx1, 4.9, color=MUTED)
icon(s, "refresh", W / 2 - 0.3, 5.15, 0.6, circle=WHITE)
tb(s, M, 6.1, W - 2 * M, 0.5, "agent 可以參與：查找、整理、試作、檢查", size=16, align="c", color=ACCENT, bold=True)
notes(s, "先看工作全貌，再談 agent 的位置。它能幫哪一段，取決於你卡在哪裡，這是下一頁。")

# ============ 03 兩種能力 ============
s = assert_slide("你卡在哪裡，由領域知識與駕馭 agent 的能力決定", "前半段")
px, py, pw, ph = picture(s, A / "week2-evidence-pilot/magnifying-glass-book.jpg", M, BODY_Y + 0.1, 5.4, 3.4)
tb(s, M, py + ph + 0.1, 5.4, 0.35, "圖片：Julo，Wikimedia Commons，公有領域，未修改。", size=10, color=MUTED)
x = 6.6; w = W - M - x
rect(s, x, BODY_Y + 0.1, w, 1.55)
tb(s, x + 0.3, BODY_Y + 0.25, w - 0.6, 0.45, "領域知識", size=19, bold=True)
tb(s, x + 0.3, BODY_Y + 0.75, w - 0.6, 0.8, "理解方法與結果的因果關係，看得出目前缺什麼。", size=14)
tb(s, x, BODY_Y + 1.65, w, 0.6, "×", size=32, bold=True, color=ACCENT, align="c", anchor="m")
rect(s, x, BODY_Y + 2.3, w, 1.55)
tb(s, x + 0.3, BODY_Y + 2.45, w - 0.6, 0.45, "駕馭 agent 的能力", size=19, bold=True)
tb(s, x + 0.3, BODY_Y + 2.95, w - 0.6, 0.8, "熟悉模型與工具的功能與限制，選得出合適的做法。", size=14)
notes(s, "例子是資料展示。回答什麼問題、用哪些資料、如何呈現，屬領域知識；知道可以請 agent 做互動網頁、怎麼試用回報，屬工具熟悉度。兩者相乘，想法才能變成可執行的要求。")

# ============ 04 四種協助 ============
s = assert_slide("缺口不同，需要的協助不同", "前半段")
helps = [("compass", "缺方向", "請它說明做法，再自行查閱。"), ("tool", "不熟工具", "描述想法請它實作，再試用修改。"),
         ("settings", "重複雜工", "交代規則與完成條件，交給它處理。"), ("zoom-check", "需要檢查", "請它找出遺漏，再自行決定是否修改。")]
cw = (W - 2 * M - 0.4) / 2
for i, (ic, h1, body) in enumerate(helps):
    x = M + (i % 2) * (cw + 0.4); y = BODY_Y + 0.1 + (i // 2) * 2.2
    card(s, x, y, cw, 2.0, h1, body, ic=ic, head_size=19, body_size=15)
notes(s, "同一件事可能同時需要幾種協助。")

# ============ 05 互動 ============
s = assert_slide("先說出一件你要向同事說明的事", "前半段")
qs = ["要說明的事是什麼？", "目前卡在哪裡？", "四種協助，想先試哪一種？"]
for i, q in enumerate(qs):
    y = BODY_Y + 0.4 + i * 1.3
    num(s, 3.4, y + 0.05, i + 1, d=0.55, size=18)
    tb(s, 4.2, y, 7, 0.65, q, size=26, anchor="m")
notes(s, "請三位學員回答，記下作為練習題目。下一步是把需要向 agent 說清楚。")

# ============ 06 請教與交辦 ============
s = assert_slide("還沒有方法時請教，已知道要做什麼時交辦", "前半段")
cols = [("message-circle", "請教", ["說出自己的想法與理由，再從它的觀點追問。", "建議須經資料或實作確認。"], "「比較幾種用 agent 做簡報的方式，比較條件該怎麼定？」"),
        ("list-check", "交辦", ["說明給誰用、做成什麼、範圍與必須保留的內容。", "做法由它提出。"], "「做成十頁左右可編輯的 PowerPoint，16:9，先試做一頁。」")]
cw = (W - 2 * M - 0.4) / 2
for i, (ic, h1, lines, ex) in enumerate(cols):
    x = M + i * (cw + 0.4)
    rect(s, x, BODY_Y + 0.1, cw, 4.3)
    icon(s, ic, x + 0.3, BODY_Y + 0.35, 0.7, circle=WHITE)
    tb(s, x + 1.15, BODY_Y + 0.4, cw - 1.4, 0.6, h1, size=22, bold=True, anchor="m")
    tb(s, x + 0.3, BODY_Y + 1.3, cw - 0.6, 1.5, [{"text": t, "bullet": True} for t in lines], size=15, spacing=8)
    tb(s, x + 0.3, BODY_Y + 3.0, cw - 0.6, 1.2, ex, size=14, bold=True, color=ACCENT)
notes(s, "兩者可來回切換；採用什麼、如何修改由自己決定。兩句例句就是示範時實際輸入的字。")

# ============ 07 對照組 ============
s = assert_slide("一句話交辦，十分鐘得到 12 頁", "前半段")
px, py, pw, ph = picture(s, A / "week2-baseline/baseline-comparison-stars.png", M, BODY_Y + 0.1, 7.6, 4.3)
tb(s, M, py + ph + 0.1, 7.6, 0.35, "課程試做，2026-09-07 由 agent 生成，未經查證。", size=10, color=MUTED)
x = 8.6; w = W - M - x
rows = [("輸入", "「幫我做一份介紹用 agent 做業務簡報有哪些方式、怎麼選的投影片，給同事看的」"), ("時間", "約十分鐘"), ("輸出", "12 頁：議程、四種方式、星等比較表、判斷問題、情境對照")]
y = BODY_Y + 0.1
for lab, val in rows:
    tb(s, x, y, w, 0.35, lab, size=12, bold=True, color=ACCENT)
    tb(s, x, y + 0.35, w, 1.1, val, size=14)
    y += 1.5
notes(s, "讓學員看一分鐘。多數人第一眼會覺得可用。")

# ============ 08 答不出為什麼 ============
s = assert_slide("這 12 頁，講的人答不出為什麼", "前半段")
px, py, pw, ph = picture(s, A / "week2-baseline/baseline-method-card.png", M, BODY_Y + 0.1, 7.6, 4.3)
tb(s, M, py + ph + 0.1, 7.6, 0.35, "對照組第 4 頁。", size=10, color=MUTED)
x = 8.6; w = W - M - x
probs = ["「資料串接 agent」是什麼，講的人自己不知道。", "星等沒有來源。", "情境全是客戶提案與品牌規範，沒有送審、列印、承辦人接手。"]
for i, t in enumerate(probs):
    y = BODY_Y + 0.15 + i * 1.4
    num(s, x, y, i + 1)
    tb(s, x + 0.6, y - 0.02, w - 0.6, 1.2, t, size=15)
notes(s, "整份通順，但沒有一頁經過自己的判斷。它好看是因為照一份文字規則做、做完看過畫面，這兩件事第六步會用到。")

# ============ 09 對照表 ============
s = assert_slide("版面規則可以交給 agent，內容的依據要自己查", "前半段")
rows = [["", "對照組", "今天要做的"], ["頁數", "12", "10 到 12"], ["版面", "依文字規則製作", "同一套規則交給 agent"],
        ["內容來源", "模型既有知識", "自己查過的材料與草稿"], ["依據", "無", "每一頁能說明理由"]]
table(s, M + 0.8, BODY_Y + 0.3, W - 2 * M - 1.6, rows, [2.2, 4.4, 4.53], size=17, row_h=0.72)
notes(s, "頁數相同、版面規則相同，差別只在內容有沒有依據。")

# ============ 10 三種形式 ============
s = assert_slide("示範題比較圖片、HTML、PPTX 三種形式", "前半段")
tb(s, M, BODY_Y - 0.15, W - 2 * M, 0.4, "用 agent 做業務簡報，該選哪種方式？", size=15, color=MUTED)
forms = [("圖片式", "image2-voltgo-city-3x2.jpg", "修改須重新產生"), ("HTML", "html-store-manager-30-60-90-3x2.jpg", "能否修改依製作方式而異"), ("PPTX", "pptx-flowpilot-2026-q2-clean-3x2.jpg", "逐頁逐字可改；送審列印沿用既有流程")]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, fn, desc) in enumerate(forms):
    x = M + i * (cw + 0.3)
    rect(s, x, BODY_Y + 0.4, cw, 3.9)
    tb(s, x + 0.25, BODY_Y + 0.55, cw - 0.5, 0.45, h1, size=20, bold=True, color=ACCENT, anchor="m")
    picture(s, A / "week2-nesa-examples" / fn, x + 0.25, BODY_Y + 1.1, cw - 0.5, 1.6)
    tb(s, x + 0.25, BODY_Y + 2.85, cw - 0.5, 1.3, desc, size=15)
tb(s, M, 6.65, W - 2 * M, 0.35, "來源：slidefirm/NESA-SLIDE README 的 demo 縮圖，MIT 授權，未修改。", size=10, color=MUTED)
notes(s, "差別在做好之後能改什麼、改了存在哪。原站將 HTML 再分兩型，本課合併。練習選預設題的人，把 repo 連結交給 agent，每類實際打開一份；只看範例，不查安裝與製作工具。")

# ============ 11 兩個迴圈 ============
s = assert_slide("有依據的內容來自兩個迴圈：先建立論點，再確認內容，最後排版", "前半段")
def fbox(x, y, w, h, text, fill=WHITE, line=LINE, size=12, bold=False, color=DARK):
    rect(s, x, y, w, h, fill=fill, line=line, radius=0.12)
    tb(s, x + 0.08, y, w - 0.16, h, text, size=size, bold=bold, color=color, align="c", anchor="m")
gy, gh, gw, bw, bh = BODY_Y + 0.15, 3.4, 5.3, 4.7, 0.66
ya, yb, yc = gy + 0.45, gy + 1.35, gy + 2.25
for gx, label, boxes, legend in [(M, "建立論點", ["帶著問題探索：蒐集文獻、案例與圖片", "人通讀資料，先寫大綱；比較依據，加入洞見與取捨", "Agent 整理故事線，人確認重點與順序"], "虛線：補探索、調整想法"),
                                  (7.0, "確認內容", ["Agent 在主草稿展開文章，沿論點完整說明", "人充實內容：改成自己的話，查核與補缺口", "整體通讀與收斂：Agent 建議刪併，人決定取捨"], "虛線：依取捨修改草稿")]:
    rect(s, gx, gy, gw, gh, fill=TINT)
    tb(s, gx + 0.2, gy + 0.08, 3, 0.3, label, size=13, bold=True, color=ACCENT)
    bx = gx + 0.3
    for k, (yy, t) in enumerate(zip((ya, yb, yc), boxes)):
        fbox(bx, yy, bw, bh, t, bold=(k == 1))
    arrow(s, bx + bw / 2, ya + bh, bx + bw / 2, yb)
    arrow(s, bx + bw / 2, yb + bh, bx + bw / 2, yc)
    arrow(s, gx + 0.16, yc + bh / 2, gx + 0.16, ya + bh / 2 + 0.05, dashed=True)
    tb(s, gx + 0.3, gy + gh - 0.32, 3, 0.3, legend, size=10, color=MUTED)
arrow(s, M + 0.3 + bw, yc + bh / 2, 7.3, ya + bh / 2, color=ACCENT, width=2)
by = gy + gh + 0.3
fbox(M, by, 5.8, 0.6, "轉成投影片分頁：選用素材，安排口述與畫面", fill=DARK, line=DARK, color=WHITE, bold=True)
fbox(7.0, by, 5.7, 0.6, "Agent 排版；人試講、檢查與微調", fill=DARK, line=DARK, color=WHITE, bold=True)
arrow(s, 7.3 + bw / 2, yc + bh, 7.0 + 5.7 / 2, by, color=ACCENT, width=2)
arrow(s, 5.8, by + 0.3, 7.0, by + 0.3, color=ACCENT, width=2)
notes(s, "人先列綱，agent 沿方向整理；agent 展開文章，人充實查核；新材料可改方向，只改受影響的段落。七步是這張圖的操作版。")

# ============ 12 保存 ============
s = assert_slide("過程保存成三類檔案，下次才能接續", "前半段")
steps = ["請 agent 寫成 Markdown", "自己開啟確認", "下次指定讀取"]
sw = 3.4
sx0 = (W - (3 * sw + 2 * 0.7)) / 2
for i, t in enumerate(steps):
    x = sx0 + i * (sw + 0.7)
    rect(s, x, BODY_Y + 0.1, sw, 0.7, fill=WHITE, line=LINE)
    num(s, x + 0.15, BODY_Y + 0.24, i + 1)
    tb(s, x + 0.7, BODY_Y + 0.1, sw - 0.8, 0.7, t, size=15, anchor="m")
    if i < 2:
        arrow(s, x + sw + 0.08, BODY_Y + 0.45, x + sw + 0.62, BODY_Y + 0.45, color=ACCENT, width=2)
folders = [("kb/", "材料與來源"), ("drafts/", "論點、文章、分頁草稿、棄案"), ("output/", "PPT 成品")]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, body) in enumerate(folders):
    x = M + i * (cw + 0.3); y = BODY_Y + 1.3
    rect(s, x, y, cw, 2.6)
    icon(s, "folder", x + cw / 2 - 0.5, y + 0.3, 1.0)
    tb(s, x, y + 1.4, cw, 0.45, h1, size=22, bold=True, align="c")
    tb(s, x + 0.3, y + 1.9, cw - 0.6, 0.6, body, size=14, color=MUTED, align="c")
notes(s, "確認時看意思是否改變、來源是否保留、未定事項是否標明。三者都在，這份簡報才接得下去。")

# ============ 13 開始前 ============
s = assert_slide("開始前先確認四件事", "示範與練習")
pre = [("shield", "資料使用設定", "關閉訓練不代表資料不上雲。"), ("lock", "repo 公開與否", "private repo 同樣在雲端。"),
       ("route", "PPTX 製作路徑", "這台機器用什麼做出可編輯的檔案，講師告知。"), ("eye", "預覽方式", "無預覽軟體時，請 agent 匯出 PDF。")]
cw = (W - 2 * M - 0.4) / 2
for i, (ic, h1, body) in enumerate(pre):
    x = M + (i % 2) * (cw + 0.4); y = BODY_Y + 0.1 + (i // 2) * 2.0
    card(s, x, y, cw, 1.8, h1, body, ic=ic, head_size=18, body_size=14)
tb(s, M, 6.45, W - 2 * M, 0.4, "接下來七頁：示範時看講師做，練習時照著做。", size=15, color=ACCENT, bold=True)
notes(s, "練習用公開或虛構材料，舊簡報檢查備註與附件。數發部手冊第 81 頁 4.3.6。示範只展開「內容常改」一條論點；學員記兩件事：自己會加哪個條件、最想先試哪一步。")

# ============ 14–20 七步：固定版式 ============
def step_slide(n, title, do, say, check, demo, prac, evidence, note):
    s = assert_slide(f"第{n}步：{title}", "示範與練習", right=f"示範 {demo} 分　練習 {prac} 分")
    lx, lw = M, 6.9
    blocks = [("做什麼", do), ("對 agent 說", say), ("確認什麼", check)]
    y = BODY_Y
    hs = [1.55, 1.25, 1.35]
    for (lab, lines), h in zip(blocks, hs):
        rect(s, lx, y, lw, h)
        tb(s, lx + 0.25, y + 0.15, 1.4, 0.35, lab, size=12, bold=True, color=ACCENT)
        if lab == "對 agent 說":
            tb(s, lx + 0.25, y + 0.5, lw - 0.5, h - 0.6, lines, size=14, bold=True, color=DARK, spacing=4)
        else:
            tb(s, lx + 0.25, y + 0.5, lw - 0.5, h - 0.6, [{"text": t, "bullet": True} for t in lines], size=14, spacing=4)
        y += h + 0.15
    rx = lx + lw + 0.4
    evidence(s, rx, BODY_Y, W - M - rx, y - 0.15 - BODY_Y)
    notes(s, note)
    return s

num_cn = ["一", "二", "三", "四", "五", "六", "七"]

def ev14(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "練習專案／", size=14, bold=True)
    tb(s, x + 0.3, y + 0.7, w - 0.6, 2.5, [{"text": "AGENTS.md", "bold": True, "after": 2}, {"text": "保存方式與共用要求", "size": 12, "color": MUTED, "after": 14},
                                          {"text": "kb/第一次討論.md", "bold": True, "after": 2}, {"text": "題目、受眾、用途、想法、未定事項", "size": 12, "color": MUTED}], size=15)
    icon(s, "folder", x + w - 1.3, y + h - 1.3, 1.0)

step_slide(num_cn[0], "建立 kb/ 與 AGENTS.md，並請 agent 讀取",
           ["開空專案，說明題目、受眾、用途。", "請 agent 建 kb/，保存方式寫進 AGENTS.md。"],
           "「我要向同事說明＿＿。先建 kb/ 保存材料與討論，保存位置和共用要求寫進 AGENTS.md。」",
           ["開啟兩個檔案，太長就刪到只留保存位置。", "再請它讀取 AGENTS.md。"], 2, 5, ev14,
           "資料夾名稱不會讓 agent 自動讀，要明確請它讀。後續依需要再建 assets/、drafts/、output/。")

def ev15(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=ACCENT)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "預設題的材料", size=14, bold=True, color=ACCENT)
    tb(s, x + 0.3, y + 0.7, w - 0.6, h - 0.9, [{"text": "把 NESA-SLIDE 連結交給 agent，每類實際開啟一份。", "bullet": True},
                                              {"text": "只看範例，不查安裝與製作工具。", "bullet": True},
                                              {"text": "PPTX 抓不到，向講師索取備份。", "bullet": True}], size=14, spacing=8)

step_slide(num_cn[1], "找材料，每份實際開啟，補上自己的條件",
           ["說受眾與情境，請 agent 找材料。", "每份實際開啟，整理進 kb/，附來源與日期。", "想到條件就補：常改、誰接手、在哪播放、送審列印。"],
           "「同事想＿＿，但不清楚有哪些方式。先幫我找材料，每一份我要實際看過。」",
           ["區分作者說明、自己看過、待確認。", "30 分停止搜尋。"], 4, 20, ev15,
           "一開始講不出條件是正常的，看過範例才會想起來。示範時材料課前查好，現場只讀取並補一條。")

def ev16(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "立場修正的例子", size=14, bold=True, color=ACCENT)
    tb(s, x + 0.3, y + 0.7, w - 0.6, 0.9, "「圖片好看但不能改」", size=15, color=MUTED)
    tb(s, x + 0.3, y + 1.35, w - 0.6, 0.4, "↓ 看過故事線", size=12, color=MUTED)
    tb(s, x + 0.3, y + 1.75, w - 0.6, 1.0, "「一次性、不再修改的內容才選圖片」", size=15, bold=True)
    tb(s, x + 0.3, y + 2.9, w - 0.6, 1.0, "棄用「互動性」比較：受眾不需要，也未實測。", size=13, color=MUTED)

step_slide(num_cn[2], "先說自己的立場，再讓 agent 整理故事線",
           ["說立場與理由，請 agent 沿方向整理故事線。", "連讀，至少修正一處。", "棄用的另存一份，記理由。"],
           "「我認為＿＿最重要，因為＿＿。照這個順序整理故事線。」",
           ["檢查點一（45 分）：kb/ 有三份材料，加一條故事線。"], 4, 15, ev16,
           "想不到時讓它先試串，自己逐段說認同或修改的理由。")

def ev17(s, x, y, w, h):
    rows = [["修改前", "HTML 改完可直接存檔，同事能繼續改"], ["對照來源", "預設只存瀏覽器本機，寫回需另備 server"],
            ["修改後", "可在瀏覽器改文字與位置；預設只存該台電腦，交接須先匯出。作者說明，未實測"]]
    table(s, x, y, w, rows, [1.2, w - 1.2], size=12, row_h=1.1)

step_slide(num_cn[3], "agent 起草，人補例子，再挑一句查證",
           ["請 agent 先寫一段，看過再展開。", "補自己的例子與經驗。", "挑一句斷言請它對照來源，前後版本都留。"],
           "「先寫一段給我看。」　「這句『＿＿』請對照來源檢查，前後版本都留在草稿裡。」",
           ["「這種方式比較好，適合所有簡報」這類結尾要補條件。", "未實測照標未實測。"], 4, 20, ev17,
           "挑一句查證的動作學員不會主動做，練習第 50 分鐘全場再示範一次。")

def ev18(s, x, y, w, h):
    segs = [("需要與條件", 2, ACCENT), ("主體", 5, DARK), ("情境", 2, MUTED), ("建議", 1, ACCENT), ("製作方式", 1, DARK), ("來源", 1, MUTED)]
    cols = 6
    cw = (w - 0.5) / cols; ch = cw * 9 / 16
    k = 0
    for name, cnt, col in segs:
        for _ in range(cnt):
            cx = x + 0.25 + (k % cols) * cw; cy = y + 0.6 + (k // cols) * (ch + 0.35)
            rect(s, cx + 0.05, cy, cw - 0.1, ch, fill=col, radius=0.05)
            tb(s, cx + 0.05, cy, cw - 0.1, ch, str(k + 1), size=11, color=WHITE, align="c", anchor="m")
            k += 1
    tb(s, x + 0.25, y + 0.15, w - 0.5, 0.35, "頁面模板：12 頁", size=14, bold=True, color=ACCENT)
    ly = y + 0.6 + 2 * (ch + 0.35) + 0.1
    tb(s, x + 0.25, ly, w - 0.5, h - (ly - y) - 0.1, [{"text": f"{n}：{name}", "size": 12} for n, name in
                                                     [("1–2", "受眾的需要與比較條件"), ("3–7", "主體，一項一頁"), ("8–9", "套入情境"), ("10", "建議與理由"), ("11", "製作方式"), ("12", "來源")]], size=12, spacing=2)

step_slide(num_cn[4], "依模板分頁，每頁對應草稿的一段",
           ["從文章整理 10 到 12 頁。", "每頁寫重點、口述兩三句、畫面、節奏一行。"],
           "「照這個模板分成十到十二頁，每頁列重點、口述兩三句、要放的畫面。」",
           ["對不上草稿的頁刪，例如 agent 自加的議程與回顧。", "第一頁保留開場。"], 2, 10, ev18,
           "口述不是逐字稿；節奏寫成「停留約 60 秒，講完第三點換頁」。")

def ev19(s, x, y, w, h):
    ph = (h - 0.9) / 2
    tb(s, x, y, w, 0.3, "首次輸出：整頁太空、字小", size=12, color=MUTED)
    picture(s, A / "week2-student-run/first-output-page1.png", x, y + 0.35, w, ph)
    tb(s, x, y + ph + 0.55, w, 0.3, "修改後：標題加大、關鍵字做成並排色塊", size=12, color=ACCENT, bold=True)
    picture(s, A / "week2-student-run/revised-page1.png", x, y + ph + 0.9, w, ph)

step_slide(num_cn[5], "先做一頁看畫面，再展開其餘頁",
           ["請 agent 做一頁，附視覺設計規則。", "匯出檢視，提具體修改。", "展開其餘頁，逐頁看。"],
           "「做成十二頁以內可編輯的 PowerPoint，16:9，先試做一頁，說明排法。版面照這份規則。」",
           ["檢查點二（95 分）：第一頁完成並看過畫面。", "每頁備註有口述與來源；文字都是文字框。"], 7, 25, ev19,
           "修改版另存，不覆蓋首次輸出。採用的規則補進 AGENTS.md。")

def ev20(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "agent 無法代做", size=14, bold=True, color=ACCENT)
    items = [("pointer", "親自點選文字，確認可編輯"), ("microphone", "實際講一頁，計時"), ("users", "請人重述重點")]
    for i, (ic, t) in enumerate(items):
        yy = y + 0.85 + i * 1.05
        icon(s, ic, x + 0.3, yy, 0.7, circle=TINT)
        tb(s, x + 1.2, yy, w - 1.5, 0.7, t, size=15, anchor="m")

step_slide(num_cn[6], "試講一頁，找出說不出理由的那頁",
           ["兩人一組，各講最沒把握的一頁，聽的人重述。", "卡住的那頁回草稿補或刪。"],
           "「估每頁口述的字數和時間。」",
           ["開頭說的需求，每項都在成品裡。", "kb/、草稿、分頁、PPTX 最新版找得到。"], 5, 10, ev20,
           "試跑時有學員開頭說長官在意圖表好不好看，成品沒處理，自己與 agent 都沒發現。")

# ============ 21 練習 ============
s = assert_slide("練習：題目自選，開始時先讓 agent 做第一版", "練習 120 分")
cols = [("tag", "題目", ["預設題可用。", "自選題：有受眾與用途、材料公開或可虛構、兩小時內找得到依據。"]),
        ("player-play", "開始時", ["題目一句話交給 agent 做第一版，背景執行。", "另開視窗走七步。"]),
        ("stack-2", "成品", ["10 到 12 頁可編輯 PPTX。", "每頁對應草稿一段。"])]
cw = (W - 2 * M - 0.6) / 3
for i, (ic, h1, lines) in enumerate(cols):
    x = M + i * (cw + 0.3)
    rect(s, x, BODY_Y + 0.1, cw, 4.0)
    icon(s, ic, x + 0.3, BODY_Y + 0.35, 0.7, circle=WHITE)
    tb(s, x + 1.15, BODY_Y + 0.4, cw - 1.4, 0.6, h1, size=21, bold=True, anchor="m")
    tb(s, x + 0.3, BODY_Y + 1.3, cw - 0.6, 2.7, [{"text": t, "bullet": True} for t in lines], size=14, spacing=8)
notes(s, "第一版留到回看。第 10 分鐘講師巡一輪，寫不出受眾與用途者改用預設題。")

# ============ 22 時間軸 ============
s = assert_slide("練習 120 分鐘分七段，兩個檢查點", "練習 120 分")
segs = [(0, 5, "開始前", ""), (5, 10, "第一步", "p14"), (10, 30, "第二步", "p15"), (30, 45, "第三步", "p16"), (45, 65, "第四步", "p17"),
        (65, 75, "第五步", "p18"), (75, 100, "第六步", "p19"), (100, 110, "第七步", "p20"), (110, 120, "回看", "p23")]
tx, tw, ty, th = M + 0.2, W - 2 * M - 0.4, 4.3, 0.55
scale = tw / 120
for i, (a, b, name, pg) in enumerate(segs):
    x = tx + a * scale; w = (b - a) * scale
    rect(s, x + 0.02, ty, w - 0.04, th, fill=(DARK if i % 2 else DARK2), radius=0.03)
    up = i % 2 == 0
    ly = ty - 0.72 if up else ty + th + 0.12
    tb(s, x - 0.3, ly, w + 0.6, 0.6, [{"text": name, "bold": True, "after": 0}, {"text": pg, "size": 10, "color": MUTED}], size=12, align="c")
for m_ in (0, 30, 45, 65, 95, 120):
    tb(s, tx + m_ * scale - 0.3, ty + th + 0.85, 0.6, 0.3, str(m_), size=10, color=MUTED, align="c")
for m_, lab in ((45, "檢查點一：三份材料與一條故事線"), (95, "檢查點二：第一頁完成並看過畫面")):
    x = tx + m_ * scale
    arrow(s, x, ty - 1.1, x, ty, color=ACCENT, width=2, head=False)
    icon(s, "flag", x - 0.05, ty - 1.55, 0.45)
    tb(s, x + 0.45, ty - 1.6, 4.2, 0.5, lab, size=13, bold=True, color=ACCENT)
tb(s, M, 6.5, W - 2 * M, 0.4, "未過檢查點：交第 2 頁、主體任一頁、建議頁，並試講一頁。", size=14, color=MUTED)
notes(s, "材料到 30 分停止搜尋。練習全程這頁與七步頁輪流投影。講師介入三處：搜尋過頭、第 50 分鐘再示範查證、首次輸出太空。")

# ============ 23 回看 ============
s = assert_slide("回看：把第一版和自己做的並排", "練習 120 分")
for i, (lab, sub) in enumerate([("第一版", "一句話生成"), ("自己做的", "七步之後")]):
    x = M + 1.0 + i * 6.4
    rect(s, x, BODY_Y + 0.1, 4.6, 2.2)
    icon(s, "stack-2", x + 0.4, BODY_Y + 0.55, 1.3)
    tb(s, x + 2.0, BODY_Y + 0.6, 2.5, 0.6, lab, size=24, bold=True, anchor="m")
    tb(s, x + 2.0, BODY_Y + 1.25, 2.5, 0.5, sub, size=14, color=MUTED)
icon(s, "arrows-left-right", W / 2 - 0.45, BODY_Y + 0.75, 0.9, circle=TINT)
qs = [("補上了什麼", "查到的材料、自己的例子、修正過的斷言"), ("哪些是自己的判斷", "採用或刪除的論點、棄案理由、開頭在意的條件"), ("哪些交給了 agent", "搜尋、整理、查證、排版")]
for i, (h1, body) in enumerate(qs):
    y = BODY_Y + 2.7 + i * 0.75
    num(s, M + 1.0, y + 0.05, i + 1, d=0.4, size=13)
    tb(s, M + 1.6, y, 3.6, 0.5, h1, size=16, bold=True, anchor="m")
    tb(s, M + 5.3, y, 7, 0.5, body, size=14, color=MUTED, anchor="m")
notes(s, "說不出理由的那頁，內容還不是自己的。補一條要求進 AGENTS.md；下次想處理的工作告訴講師。")

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
