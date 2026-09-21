"""第四週 4-1 的圖：以 python-pptx 原生圖形畫在投影片上，可在 PowerPoint 直接改。

每個圖函式的簽名都是 draw(slide, x, y, w, h)，單位英吋，畫在指定區域內。
build-week4.py 依頁標題查 FIGURES 取函式。
"""
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.dml import MSO_LINE
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

DARK = RGBColor(0x2A, 0x26, 0x23)
MUTED = RGBColor(0x6B, 0x65, 0x60)
ACCENT = RGBColor(0xB8, 0x50, 0x42)
TINT = RGBColor(0xEE, 0xF2, 0xEF)
LINE = RGBColor(0xD9, 0xD4, 0xCE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AGENT = RGBColor(0x3F, 0x7D, 0x7A)      # Agent 與模型
AGENT_LT = RGBColor(0xD6, 0xE6, 0xE4)
PROG = RGBColor(0x5B, 0x6B, 0x8C)       # 程式、工具、流程節點
PROG_LT = RGBColor(0xDD, 0xE2, 0xEC)
FADE = RGBColor(0xC9, 0xC4, 0xBE)
FADE_LT = RGBColor(0xF4, 0xF2, 0xEF)
FONT = "PingFang TC"


# ---------- 基本圖形 ----------
def _txt(tf, text, size, bold=False, color=DARK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    lines = text.split("\n") if text else [""]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color


def box(slide, x, y, w, h, text="", fill=WHITE, line=LINE, size=14, bold=False, color=DARK,
        shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=False, lw=1.25, align=PP_ALIGN.CENTER):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
        if dash:
            s.line.dash_style = MSO_LINE.DASH
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = 0.18
        except Exception:
            pass
    _txt(s.text_frame, text, size, bold, color, align)
    return s


def text(slide, x, y, w, h, s, size=12, color=MUTED, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _txt(tb.text_frame, s, size, bold, color, align, anchor)
    return tb


def arrow(slide, x1, y1, x2, y2, color=MUTED, dashed=False, lw=1.5, head=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(lw)
    if dashed:
        c.line.dash_style = MSO_LINE.DASH
    if head:
        ln = c.line._get_or_add_ln()
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
    return c


def person(slide, cx, cy, size=0.36, color=ACCENT):
    """小人：圓頭加圓肩。cx, cy 為圖示中心。"""
    r = size * 0.36
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r / 2), Inches(cy - size / 2), Inches(r), Inches(r))
    body = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx - size / 2), Inches(cy - size / 2 + r * 1.05),
                                  Inches(size), Inches(size - r * 1.05))
    for s in (head, body):
        s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background(); s.shadow.inherit = False
    try:
        body.adjustments[0] = 0.5
    except Exception:
        pass
    return head, body


def doc(slide, x, y, w, h, label="", fill=WHITE, line=MUTED, size=11, color=DARK):
    return box(slide, x, y, w, h, label, fill=fill, line=line, size=size, color=color, shape=MSO_SHAPE.FOLDED_CORNER)


def chip(slide, x, y, w, h, label, fill=AGENT, color=WHITE, size=11, bold=True):
    return box(slide, x, y, w, h, label, fill=fill, line=None, size=size, bold=bold, color=color)


# ---------- 圖一：五格流程圖（第 3、18 頁） ----------
STEPS = [
    ("讀取", "挑事件\n取第一報與地震報告"),
    ("計算", "報時差、震央距離差、規模差\n分陸上海域"),
    ("確認", "納入、取哪一報\n異常解釋"),
    ("產出", "統計表、分布圖、年報段落\n送審"),
    ("保存", "規則、事件清單\n腳本回存"),
]
BADGES = ["2 或 5", "3", "人；決行\n是否納入可交 4", "2\n表圖與段落", "3\n規則進腳本"]


def fig_five_steps(slide, x, y, w, h, badges=False):
    n = len(STEPS)
    gap = 0.45
    bw = (w - gap * (n - 1)) / n
    bh = 1.25
    by = y + (h - bh) / 2 + 0.15
    for i, (name, sub) in enumerate(STEPS):
        bx = x + i * (bw + gap)
        human = name in ("確認", "產出")
        box(slide, bx, by, bw, bh, "", fill=TINT if human else PROG_LT, line=ACCENT if human else PROG, lw=1.5)
        text(slide, bx, by + 0.08, bw, 0.4, name, size=16, bold=True, color=DARK)
        text(slide, bx + 0.05, by + 0.48, bw - 0.1, 0.75, sub, size=10.5, color=MUTED)
        if human:
            person(slide, bx + bw - 0.28, by + 0.26, size=0.3)
        if i < n - 1:
            arrow(slide, bx + bw + 0.03, by + bh / 2, bx + bw + gap - 0.03, by + bh / 2, color=MUTED, lw=1.75)
        if badges:
            chip(slide, bx + 0.05, by - 0.62, bw - 0.1, 0.52, BADGES[i], fill=ACCENT if not human else DARK, size=11)
    # 材料在讀取上方、成品在產出下方
    if not badges:
        mx = x
        doc(slide, mx, y + 0.05, bw, 0.62, "地震目錄、報次紀錄\n去年的表", fill=WHITE, line=MUTED, size=10.5)
        arrow(slide, mx + bw / 2, y + 0.69, mx + bw / 2, by - 0.03, color=MUTED, lw=1.25)
    px = x + 3 * (bw + gap)
    doc(slide, px, by + bh + 0.45, bw, 0.6, "年報的表與圖", fill=WHITE, line=ACCENT, size=11)
    arrow(slide, px + bw / 2, by + bh + 0.03, px + bw / 2, by + bh + 0.42, color=MUTED, lw=1.25)


def fig_five_steps_badges(slide, x, y, w, h):
    fig_five_steps(slide, x, y, w, h, badges=True)


# ---------- 圖二：模型與執行框架（第 4 頁） ----------
def fig_model_vs_harness(slide, x, y, w, h):
    mid = x + w / 2
    arrow(slide, mid, y, mid, y + h, color=LINE, lw=1, head=False)
    # 左：單獨的模型
    lx, lw_ = x, w / 2 - 0.4
    text(slide, lx, y, lw_, 0.35, "單獨的模型", size=13, bold=True, color=DARK)
    m = box(slide, lx + 0.3, y + 0.9, 1.7, 1.0, "模型", fill=AGENT, line=None, size=16, bold=True, color=WHITE)
    arrow(slide, lx + 2.05, y + 1.4, lx + 2.75, y + 1.4, color=MUTED)
    doc(slide, lx + 2.8, y + 0.85, 1.9, 1.1, "文字\n（計算的程式碼）", fill=WHITE, line=DARK, size=11)
    # 灰的工具，沒有連線
    box(slide, lx + 0.4, y + 2.55, 1.4, 0.55, "檔案", fill=FADE_LT, line=FADE, size=11, color=FADE)
    box(slide, lx + 2.1, y + 2.55, 1.4, 0.55, "程式", fill=FADE_LT, line=FADE, size=11, color=FADE)
    box(slide, lx + 3.8, y + 2.55, 1.4, 0.55, "網路", fill=FADE_LT, line=FADE, size=11, color=FADE)
    text(slide, lx, y + 3.25, lw_, 0.35, "數字尚未算出；工具沒有連線", size=11, color=MUTED)
    # 右：接上執行框架
    rx, rw = mid + 0.4, w / 2 - 0.4
    text(slide, rx, y, rw, 0.35, "接上執行框架", size=13, bold=True, color=DARK)
    box(slide, rx + 0.1, y + 0.9, 1.5, 1.0, "模型", fill=AGENT, line=None, size=16, bold=True, color=WHITE)
    arrow(slide, rx + 1.65, y + 1.25, rx + 2.25, y + 1.25, color=MUTED)
    text(slide, rx + 1.55, y + 0.85, 0.8, 0.35, "呼叫", size=9.5, color=MUTED)
    box(slide, rx + 2.3, y + 0.75, 1.6, 1.3, "執行框架", fill=PROG, line=None, size=14, bold=True, color=WHITE)
    arrow(slide, rx + 2.25, y + 1.6, rx + 1.65, y + 1.6, color=MUTED)
    text(slide, rx + 1.55, y + 1.62, 0.8, 0.35, "回傳", size=9.5, color=MUTED)
    tools = ["檔案", "程式", "網路"]
    for i, t in enumerate(tools):
        ty = y + 0.55 + i * 0.75
        box(slide, rx + 4.6, ty, 1.3, 0.55, t, fill=PROG_LT, line=PROG, size=11)
        arrow(slide, rx + 3.95, y + 1.4, rx + 4.55, ty + 0.27, color=MUTED, lw=1.0)
    text(slide, rx, y + 3.25, rw, 0.35, "動作才發生：寫入磁碟、跑程式、回讀結果", size=11, color=MUTED)


# ---------- 圖三：六格示意圖（第 6 到 13 頁） ----------
SIX = ["對話", "交辦", "製作工具", "結構化判斷", "流程指派 Agent", "Agent 管理流程"]
SIX_TAGS = ["對話式", "Agent", "Agent，vibe coding", "API，決策模型", "API，常駐 Agent", "API，多 Agent 編排"]


def _flow(slide, x, y, w, faded, n=3, hole=None, fill=None, sc=1.0):
    """底部的小流程線：n 個小方塊。hole 指定某格留空（由呼叫者放別的東西）。回傳各格中心 x。"""
    bw, bh = 0.42 * sc, 0.3 * sc
    gap = (w - n * bw) / (n - 1)
    cs = []
    col = FADE if faded else PROG
    for i in range(n):
        bx = x + i * (bw + gap)
        cs.append(bx + bw / 2)
        if hole == i:
            continue
        box(slide, bx, y, bw, bh, "", fill=(FADE_LT if faded else (fill or PROG_LT)), line=col, lw=1.0, shape=MSO_SHAPE.RECTANGLE)
    for i in range(n - 1):
        arrow(slide, cs[i] + bw / 2 + 0.02, y + bh / 2, cs[i + 1] - bw / 2 - 0.02, y + bh / 2, color=col, lw=1.0, head=False)
    return cs, bh


def _cell(slide, k, x, y, w, h, active, faded, tag=None):
    """一格：以 1.95 吋高為基準設計，依實際高度等比縮放。"""
    sc = min(1.0, h / 1.95)
    def X(v): return x + v * sc
    def Y(v): return y + v * sc
    def S(v): return v * sc
    def F(v): return max(7.5, v * sc)
    fill = TINT if active else WHITE
    line = ACCENT if active else LINE
    box(slide, x, y, w, h, "", fill=fill, line=line, lw=2.0 if active else 1.0)
    tcol = FADE if faded else DARK
    text(slide, X(0.12), Y(0.06), S(0.4), S(0.3), str(k + 1), size=F(11), bold=True,
         color=ACCENT if active else (FADE if faded else MUTED), align=PP_ALIGN.LEFT)
    text(slide, x, y + h - S(0.42), w, S(0.36), SIX[k], size=F(14), bold=True, color=tcol)
    if tag:
        chip(slide, x + w - S(1.85), Y(0.1), S(1.75), S(0.3), tag, fill=ACCENT, size=F(9.5))
    ix, iy = X(0.25), Y(0.45)
    iw, ih = w - S(0.5), h - S(0.95)
    ac = FADE if faded else AGENT
    pc = FADE if faded else ACCENT
    mc = FADE if faded else MUTED
    fy = iy + ih - S(0.3)
    if k == 0:   # 對話：人與 Agent 來回，流程在旁邊沒有連線
        person(slide, ix + S(0.4), iy + S(0.45), size=S(0.4), color=pc)
        chip(slide, ix + S(1.0), iy + S(0.25), S(0.9), S(0.4), "Agent", fill=ac, size=F(10))
        arrow(slide, ix + S(0.7), iy + S(0.38), ix + S(0.97), iy + S(0.38), color=mc, lw=1.0)
        arrow(slide, ix + S(0.97), iy + S(0.52), ix + S(0.7), iy + S(0.52), color=mc, lw=1.0)
        _flow(slide, ix + iw - S(1.6), fy, S(1.5), True, sc=sc)
    elif k == 1:  # 交辦：人 → Agent → 檔案，Agent 虛線退出
        person(slide, ix + S(0.3), iy + S(0.45), size=S(0.4), color=pc)
        arrow(slide, ix + S(0.55), iy + S(0.45), ix + S(0.85), iy + S(0.45), color=mc, lw=1.0)
        chip(slide, ix + S(0.9), iy + S(0.25), S(0.9), S(0.4), "Agent", fill=ac, size=F(10))
        arrow(slide, ix + S(1.85), iy + S(0.45), ix + S(2.15), iy + S(0.45), color=mc, lw=1.0)
        doc(slide, ix + S(2.2), iy + S(0.2), S(0.7), S(0.5), "檔案", fill=WHITE, line=mc, size=F(9.5), color=tcol)
        arrow(slide, ix + S(1.35), iy + S(0.7), ix + S(1.35), iy + S(1.1), color=mc, dashed=True, lw=1.0)
        text(slide, ix + S(1.4), iy + S(0.78), S(0.6), S(0.3), "退出", size=F(9), color=mc, align=PP_ALIGN.LEFT)
    elif k == 2:  # 製作工具：人與 Agent 寫出程式，程式接進流程，Agent 退出
        person(slide, ix + S(0.3), iy + S(0.35), size=S(0.36), color=pc)
        chip(slide, ix + S(0.65), iy + S(0.17), S(0.85), S(0.36), "Agent", fill=ac, size=F(10))
        cs, bh = _flow(slide, ix + S(0.2), fy, iw - S(0.4), faded, hole=1, sc=sc)
        box(slide, cs[1] - S(0.4), fy - S(0.05), S(0.8), S(0.4), "程式", fill=(FADE_LT if faded else PROG), line=None,
            size=F(9.5), bold=True, color=WHITE if not faded else FADE, shape=MSO_SHAPE.RECTANGLE)
        arrow(slide, ix + S(1.05), iy + S(0.55), cs[1], fy - S(0.08), color=mc, lw=1.0)
        arrow(slide, ix + S(1.55), iy + S(0.35), ix + S(2.2), iy + S(0.35), color=mc, dashed=True, lw=1.0)
        text(slide, ix + S(2.2), iy + S(0.2), S(0.6), S(0.3), "退出", size=F(9), color=mc, align=PP_ALIGN.LEFT)
    elif k == 3:  # 結構化判斷：流程節點問模型一句，模型回選項
        cs, bh = _flow(slide, ix + S(0.2), fy, iw - S(0.4), faded, sc=sc)
        box(slide, cs[1] - S(0.5), iy + S(0.05), S(1.0), S(0.45), "模型", fill=ac, line=None, size=F(10), bold=True, color=WHITE)
        arrow(slide, cs[1] - S(0.12), fy - S(0.02), cs[1] - S(0.12), iy + S(0.53), color=mc, lw=1.0)
        arrow(slide, cs[1] + S(0.12), iy + S(0.53), cs[1] + S(0.12), fy - S(0.02), color=mc, lw=1.0)
        text(slide, cs[1] - S(1.4), iy + S(0.55), S(1.2), S(0.3), "狀態與問題", size=F(8.5), color=mc, align=PP_ALIGN.RIGHT)
        text(slide, cs[1] + S(0.2), iy + S(0.55), S(1.2), S(0.3), "選項、是非", size=F(8.5), color=mc, align=PP_ALIGN.LEFT)
    elif k == 4:  # 流程指派 Agent：事件觸發，Agent 迴圈用工具，交下一節點
        cs, bh = _flow(slide, ix + S(0.2), fy, iw - S(0.4), faded, hole=1, sc=sc)
        chip(slide, cs[1] - S(0.45), fy - S(0.1), S(0.9), S(0.5), "Agent", fill=ac, size=F(10))
        loop = slide.shapes.add_shape(MSO_SHAPE.CIRCULAR_ARROW, Inches(cs[1] - S(0.3)), Inches(iy + S(0.05)), Inches(S(0.6)), Inches(S(0.6)))
        loop.fill.solid(); loop.fill.fore_color.rgb = ac; loop.line.fill.background(); loop.shadow.inherit = False
        box(slide, cs[1] + S(0.5), iy + S(0.12), S(0.75), S(0.42), "工具", fill=(FADE_LT if faded else PROG_LT),
            line=(FADE if faded else PROG), size=F(9.5), color=tcol)
        text(slide, cs[0] - S(0.6), fy - S(0.42), S(1.2), S(0.3), "事件", size=F(8.5), color=mc)
    elif k == 5:  # Agent 管理流程：主 Agent 分派，在核准點停下
        cs, bh = _flow(slide, ix + S(0.2), fy, iw - S(0.4), faded, sc=sc)
        chip(slide, ix + iw / 2 - S(0.55), iy + S(0.05), S(1.1), S(0.45), "主 Agent", fill=ac, size=F(10))
        for c in cs:
            arrow(slide, ix + iw / 2, iy + S(0.52), c, fy - S(0.03), color=mc, lw=1.0)
        person(slide, cs[1] + S(0.55), fy + S(0.15), size=S(0.3), color=pc)
        text(slide, cs[1] + S(0.7), fy - S(0.02), S(0.7), S(0.3), "核准", size=F(8.5), color=mc, align=PP_ALIGN.LEFT)


def _six_grid(slide, x, y, w, h, active=None, tags=None, bottom=0.0):
    cols, rows = 3, 2
    gx, gy = 0.3, 0.25
    cw = (w - gx * (cols - 1)) / cols
    ch = (h - bottom - gy * (rows - 1)) / rows
    cells = []
    for k in range(6):
        r, c = divmod(k, cols)
        cx, cy = x + c * (cw + gx), y + r * (ch + gy)
        is_active = (active is None) or (k in active)
        faded = active is not None and k not in active
        _cell(slide, k, cx, cy, cw, ch, is_active and active is not None, faded, (tags or {}).get(k))
        cells.append((cx, cy, cw, ch))
    return cells


def fig_six(slide, x, y, w, h):
    _six_grid(slide, x, y, w, h)


def _mk_six_active(k):
    def f(slide, x, y, w, h):
        _six_grid(slide, x, y, w, h, active={k}, tags={k: SIX_TAGS[k]})
    return f


def fig_six_summary(slide, x, y, w, h):
    tags = {2: "已在改變大眾做法", 3: "探索中", 4: "探索中", 5: "探索中"}
    _six_grid(slide, x, y, w, h, active={2}, tags=tags, bottom=1.05)
    by = y + h - 0.72
    arrow(slide, x + 2.6, by, x + w - 0.2, by, color=ACCENT, lw=2.0)
    text(slide, x, by - 0.16, 2.5, 0.3, "交給 Agent 的信任", size=10.5, color=ACCENT, bold=True, align=PP_ALIGN.RIGHT)
    arrow(slide, x + 2.6, by + 0.45, x + w - 0.2, by + 0.45, color=PROG, lw=2.0)
    text(slide, x, by + 0.29, 2.5, 0.3, "成本、採用的技術難度", size=10.5, color=PROG, bold=True, align=PP_ALIGN.RIGHT)


# ---------- 圖四：七站流程線（第 14 到 17 頁） ----------
STATIONS = ["備料", "會辦", "擬稿", "執行", "審核", "決行", "歸檔"]
ROLES = ["整理者", "討論者、對手", "起草者", "操作員", "審查者", "人", "記錄者"]
DARK_STATIONS = {"擬稿", "執行", "歸檔"}


def _stations(slide, x, y, w, h, shade=False, roles=False, tag_decision=None, extra_circles=False):
    n = len(STATIONS)
    gap = 0.32
    bw = (w - gap * (n - 1)) / n
    bh = 1.05
    by = y + 1.1
    centers = []
    for i, name in enumerate(STATIONS):
        bx = x + i * (bw + gap)
        centers.append(bx + bw / 2)
        human = name == "決行"
        if human:
            fill, line, tcol = TINT, ACCENT, DARK
        elif shade:
            dark = name in DARK_STATIONS
            fill, line, tcol = (PROG, None, WHITE) if dark else (PROG_LT, PROG, DARK)
        else:
            fill, line, tcol = WHITE, PROG, DARK
        box(slide, bx, by, bw, bh, "", fill=fill, line=line, lw=1.5)
        if human:
            person(slide, bx + bw / 2, by + 0.38, size=0.38)
        else:
            doc(slide, bx + bw / 2 - 0.2, by + 0.14, 0.4, 0.42, "", fill=WHITE if not (shade and name in DARK_STATIONS) else PROG_LT, line=MUTED, size=8)
        text(slide, bx, by + 0.62, bw, 0.36, name, size=14, bold=True, color=tcol)
        if i < n - 1:
            arrow(slide, bx + bw + 0.03, by + bh / 2, bx + bw + gap - 0.03, by + bh / 2, color=MUTED, lw=1.5)
        if roles:
            rc = ACCENT if human else (WHITE if (shade and name in DARK_STATIONS) else DARK)
            rf = None if human else ((PROG if shade and name in DARK_STATIONS else PROG_LT) if shade else None)
            if rf:
                chip(slide, bx, by + bh + 0.18, bw, 0.4, ROLES[i], fill=rf, color=rc, size=11)
            else:
                text(slide, bx, by + bh + 0.18, bw, 0.4, ROLES[i], size=12, bold=True, color=rc)
    # 追蹤迴圈：歸檔上方回到備料
    ly = by - 0.45
    arrow(slide, centers[-1], by - 0.03, centers[-1], ly, color=MUTED, lw=1.25, head=False)
    arrow(slide, centers[-1], ly, centers[0], ly, color=MUTED, lw=1.25, head=False)
    arrow(slide, centers[0], ly, centers[0], by - 0.05, color=MUTED, lw=1.25)
    text(slide, x + w / 2 - 1.0, ly - 0.32, 2.0, 0.3, "追蹤" + ("：協調者" if roles else ""), size=11, color=MUTED, bold=roles)
    if tag_decision:
        chip(slide, centers[5] - 0.55, by - 0.5, 1.1, 0.36, tag_decision, fill=ACCENT, size=11)
    if extra_circles:
        for j, lab in enumerate(["對外發言", "帶動團隊"]):
            cx = x + w - 3.4 + j * 1.7
            c = box(slide, cx, y - 0.35, 1.35, 0.75, lab, fill=WHITE, line=ACCENT, size=12, bold=True, color=ACCENT, shape=MSO_SHAPE.OVAL, dash=True)
        chip(slide, centers[5] - 0.55, by - 0.5, 1.1, 0.36, "責任", fill=ACCENT, size=11)
    if shade and roles:
        text(slide, x, by + bh + 0.7, w, 0.3, "淺色：只讀　深色：寫檔、執行、寫筆記", size=10.5, color=MUTED)


def fig_stations(slide, x, y, w, h):
    _stations(slide, x, y, w, h)


def fig_stations_raci(slide, x, y, w, h):
    _stations(slide, x, y, w, h, tag_decision="當責")


def fig_stations_roles(slide, x, y, w, h):
    _stations(slide, x, y, w, h, shade=True, roles=True)


def fig_stations_outside(slide, x, y, w, h):
    _stations(slide, x, y, w, h, extra_circles=True)


FIGURES = {
    "示範：拆解一件工作": fig_five_steps,
    "模型與執行框架": fig_model_vs_harness,
    "六種介入方式": fig_six,
    "對話": _mk_six_active(0),
    "交辦": _mk_six_active(1),
    "製作工具": _mk_six_active(2),
    "結構化判斷": _mk_six_active(3),
    "流程指派 Agent": _mk_six_active(4),
    "Agent 管理流程": _mk_six_active(5),
    "總結：取捨與現況": fig_six_summary,
    "角色：Agent 是幕僚": fig_stations,
    "先畫責任線": fig_stations_raci,
    "沿流程派角色": fig_stations_roles,
    "站不進去的位置": fig_stations_outside,
    "套回案例：一件工作用幾種": fig_five_steps_badges,
}
