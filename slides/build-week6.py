#!/usr/bin/env python3
"""第六週投影片建置：把 drafts/18-week6-slides.md 產成素版 PPTX（slides/week6.pptx）。

用法：python3 slides/build-week6.py

配色、封面、段落頁、圖片處理沿用 build-week5.py，另加：
- 段落：畫面中「- 段落：……」排成不帶項目符號的內文，用於一小段完整的敘述。
- 引文：「- 引文：「原話」（日期）」排成淡底引文框，出處以小字標示。
- 表格：畫面中的 Markdown 表格照原樣畫出。
- 版面依內容決定，也可用「- 版面：主張／圖為主／左文右圖／文字」指定：
  主張＝一句放大的主張加少量補充；圖為主＝文字在上、圖在下占滿寬；
  左文右圖＝文字多時圖縮小；文字＝無圖，全寬排列。
素版原則：只做文字階層與圖片位置，不做視覺設計；圖片檔不存在時以灰框標示預留位置。
"""
import importlib.util
import os
import re
import sys
from pathlib import Path

from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("bw5", HERE / "build-week5.py")
bw5 = importlib.util.module_from_spec(spec)
sys.path.insert(0, str(HERE))
spec.loader.exec_module(bw5)

bw5.SRC = bw5.ROOT / "drafts/18-week6-slides.md"
bw5.OUT = HERE / "week6.pptx"
bw4 = bw5.bw4
W, H, M = bw5.W, bw5.H, bw5.M
PARA = "段落："


LAYOUT = "版面："


def draft_images():
    """講師本機的草稿截圖對照表（環境變數 WEEK6_DRAFT_IMAGES 指向的 TSV，不進 repo）。
    每列「頁碼<TAB>圖片路徑<TAB>圖說」，路徑相對於該檔；有對照的頁以這些截圖取代預留位置。"""
    f = os.environ.get("WEEK6_DRAFT_IMAGES")
    out = {}
    if not f or not Path(f).exists():
        return out
    base = Path(f).parent
    for line in Path(f).read_text(encoding="utf-8").splitlines():
        if not line.strip() or line.startswith("#"):
            continue
        n, rel, cap = (line.split("\t") + ["", ""])[:3]
        out.setdefault(int(n), []).append((str(base / rel), f"{cap}【草稿截圖】" if cap else ""))
    return out


DRAFT = draft_images()


def pic_at(slide, rel, box):
    """自繪圖不加框，截圖加細框。"""
    x, y, w, h = box
    path = bw5.ROOT / rel
    if not path.exists():
        bw5.placeholder(slide, x, y, w, h, rel)
        return
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if "-diagrams/" not in rel:
        pic.line.color.rgb = bw5.GREY
        pic.line.width = Pt(0.75)


bw5.pic_at = pic_at
QUOTE = "引文："


def para(slide, x, y, w, text, size, color=None, bold=False, spacing=1.15):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(1))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.line_spacing = spacing
    r = p.add_run(); r.text = text
    bw4._run(r, size, bold, color or bw5.DARK)
    h = bw5.text_height([text], size, w) * spacing * 0.98
    box.height = Inches(h + 0.1)
    return h


def quote(slide, x, y, w, text, size=17):
    """使用者原話：淡底、左側色條，出處以小字附在後面。"""
    m = re.match(r"^「?(.+?)」?(（[^）]*）)?$", text)
    body, src = m.group(1), (m.group(2) or "").strip("（）")
    th = bw5.text_height([body], size, w - 0.7) * 1.1
    h = th + (0.45 if src else 0.25)
    bw5.rect(slide, x, y, w, h, fill=bw5.IVORY)
    bw5.rect(slide, x, y, 0.08, h, fill=bw5.ACCENT)
    para(slide, x + 0.35, y + 0.12, w - 0.6, f"「{body}」", size, spacing=1.1)
    if src:
        bw5.textbox(slide, x + 0.35, y + h - 0.36, w - 0.6, 0.3, [f"開發者，{src}"], size=11,
                    color=bw5.MUTED, spacing=0)
    return h


def flow(slide, x, y, w, blocks, scale=1.0):
    """依出現順序排版：段落、引文、「**關鍵詞**：說明」、條列、表格。回傳結束的 y。"""
    ps, ks, bs = 15 * scale, 16 * scale, 15 * scale
    for kind, val in blocks:
        if kind == "table":
            nrow = len(val)
            size = 14 if nrow <= 5 else 13
            y += bw4.table(slide, x, y, w, val, size=size, row_h=0.46) + 0.25
            continue
        if kind == "num":
            y += para(slide, x, y, w, val, bs) + 0.1
            continue
        if val.startswith(PARA):
            y += para(slide, x, y, w, val[len(PARA):].strip(), ps) + 0.18
            continue
        if val.startswith(QUOTE):
            y += quote(slide, x, y, w, val[len(QUOTE):].strip()) + 0.22
            continue
        m = re.match(r"^\*\*(.+?)\*\*[：:]\s*(.*)$", val)
        if m:
            y += para(slide, x, y, w, m.group(1), ks, bold=True) + 0.02
            if m.group(2):
                y += para(slide, x, y, w, m.group(2), bs - 1, color=bw5.MUTED) + 0.14
            continue
        y += para(slide, x, y, w, "•  " + val, bs) + 0.08
    return y


def row(slide, imgs, x, y, w, h):
    """圖為主：所有圖排成一列、各自置中於等寬的格子，圖說在各圖下方。"""
    gap = 0.3
    ch = bw5.CAP_H if any(c for _, c in imgs) else 0
    cw = (w - gap * (len(imgs) - 1)) / len(imgs)
    for i, (rel, cap) in enumerate(imgs):
        cx = x + i * (cw + gap)
        bx, by, bwid, bh = bw5.fit_box(rel, cx, y, cw, h - ch)
        bx = cx + (cw - bwid) / 2
        by = y
        bw5.pic_at(slide, rel, (bx, by, bwid, bh))
        if cap:
            bw5.textbox(slide, cx, by + bh + 0.05, cw, ch - 0.05, [cap],
                        size=10.5, color=bw5.MUTED, spacing=0, align="c")


def content(prs, page, total, section=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bw5.rect(s, 0, 0, W, H, fill=bw5.WHITE)
    imgs = bw5.images_of(page)
    if page["n"] in DRAFT and (not imgs or any(not (bw5.ROOT / r).exists() for r, _ in imgs)):
        imgs = DRAFT[page["n"]]
    blocks = [(k, v) for k, v in page["blocks"] if not (k == "bullet" and v.startswith(LAYOUT))]
    hint = next((v[len(LAYOUT):].strip() for k, v in page["blocks"]
                 if k == "bullet" and v.startswith(LAYOUT)), "")
    lead = [v for k, v in blocks if k == "lead"]
    body = [(k, v) for k, v in blocks if k != "lead"]
    chars = sum(len(v) if isinstance(v, str) else sum(len(c) for r in v for c in r) for _, v in body)
    full = W - 2 * M
    bottom = H - 0.5
    if section:
        bw5.textbox(s, M, 0.38, full, 0.35, [section], size=12, color=bw5.MUTED)

    if not hint:
        if imgs:
            hint = "圖為主" if chars <= 130 else "左文右圖"
        elif chars <= 110 and all(k != "table" for k, _ in body):
            hint = "主張"
        else:
            hint = "文字"

    y = 0.78
    if hint == "主張":
        # 一句主張放大置中，少量補充在下
        y = 1.5
        bw5.textbox(s, M, y, full, 0.5, [page["title"]], size=18, bold=True, color=bw5.MUTED)
        y += 0.65
        bw5.rect(s, M, y, 0.9, 0.05, fill=bw5.ACCENT)
        y += 0.35
        for t in lead:
            y += para(s, M, y, full * 0.85, t, 32, bold=True, spacing=1.2) + 0.35
        y = flow(s, M, y, full * 0.8, body, scale=1.2)
    else:
        ch = bw5.text_height([page["title"]], 30, full)
        bw5.textbox(s, M, y, full, ch, [page["title"]], size=30, bold=True, color=bw5.DARK)
        y += ch + 0.12
        bw5.rect(s, M, y, 0.9, 0.05, fill=bw5.ACCENT)
        y += 0.4
        col = full
        if hint == "左文右圖":
            ratio = 0.36 if chars > 260 else 0.45
            vx = M + full * (1 - ratio) + 0.2
            boxes, left = bw5.visual_boxes(imgs, vx, y, W - M - vx, bottom - y)
            for rel, cap, box in boxes:
                bw5.pic_at(s, rel, box)
                if cap:
                    bx, by, bwid, bh = box
                    bw5.textbox(s, bx, by + bh + 0.05, max(bwid, 2.2), bw5.CAP_H - 0.05, [cap],
                                size=10.5, color=bw5.MUTED, spacing=0)
            col = left - 0.35 - M
        for t in lead:
            y += para(s, M, y, col, t, 20, bold=True) + 0.22
        y = flow(s, M, y, col, body)
        if hint == "圖為主":
            y += 0.1
            row(s, imgs, M, y, full, bottom - y)
    if y > bottom + 0.1:
        print(f"  [溢出警告] 第 {page['n']} 頁「{page['title']}」（{hint}）估計高度到 {y:.1f} in")
    bw5.textbox(s, W - M - 0.8, H - 0.45, 0.8, 0.3, [f"{page['n']} / {total}"], size=10,
                color=bw5.MUTED, align="c")
    bw5.notes(s, page["notes"])


bw5.content = content

_parse = bw4.parse


def parse(md_path):
    """逐頁稿結尾的「待準備」是製作備註，不進最後一頁的備註。"""
    pages = _parse(md_path)
    last = pages[-1]["notes"]
    if "---" in last:
        del last[last.index("---"):]
    return pages


bw4.parse = parse

if __name__ == "__main__":
    bw5.main()
