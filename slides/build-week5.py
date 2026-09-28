#!/usr/bin/env python3
"""第五週投影片建置：把 drafts/15-week5-slides.md 產成素版 PPTX（slides/week5.pptx）。

用法：python3 slides/build-week5.py

沿用 build-week4.py 的逐頁稿解析與文字元件，另加：
- 圖片：畫面中「- 圖：`路徑`」（PNG、JPG）依序置入；路徑不存在時以灰框標示預留位置與檔名，供使用者之後替換。
- 版面：標題、主訊息、條列在上，圖片並排在下；文字較多時改為左文右圖。
- 封面：若有圖，滿版鋪底，標題疊在左側留白處。
- 配色：沿用 assets/week5-visuals/README.md 的封面規格（象牙米白、暖灰、陶土紅、灰藍）。

素版原則：只做文字階層與圖片位置，不做視覺設計，供使用者套用素材。
"""
import importlib.util
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
spec = importlib.util.spec_from_file_location("bw4", HERE / "build-week4.py")
bw4 = importlib.util.module_from_spec(spec)
sys.path.insert(0, str(HERE))
spec.loader.exec_module(bw4)

SRC = ROOT / "drafts/15-week5-slides.md"
OUT = HERE / "week5.pptx"

DARK = RGBColor(0x2A, 0x26, 0x23)
MUTED = RGBColor(0x6B, 0x65, 0x60)
ACCENT = RGBColor(0xBE, 0x89, 0x79)   # 陶土紅
BLUE = RGBColor(0xA6, 0xB7, 0xBA)     # 灰藍
IVORY = RGBColor(0xF5, 0xF2, 0xEC)    # 象牙米白
GREY = RGBColor(0xD8, 0xD2, 0xC9)     # 暖灰
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W, H, M = bw4.W, bw4.H, bw4.M

textbox = bw4.textbox
notes = bw4.notes


def rect(slide, x, y, w, h, fill, line=None):
    return bw4.rect(slide, x, y, w, h, fill=fill, line=line)


def images_of(page):
    """畫面中「圖：`路徑`」進了備註；依序取出路徑，保留不存在者作為預留位置。"""
    out = []
    for line in page["notes"]:
        if line.startswith("圖："):
            p = line[len("圖："):].strip().strip("`")
            if p and p not in out:
                out.append(p)
    page["notes"] = [l for l in page["notes"] if not l.startswith("圖：")]
    return out


def placeholder(slide, x, y, w, h, rel):
    rect(slide, x, y, w, h, fill=IVORY, line=GREY)
    textbox(slide, x + 0.1, y + h / 2 - 0.3, w - 0.2, 0.6, ["圖片預留", Path(rel).name],
            size=11, color=MUTED, align="c", anchor="m", spacing=2)


def place(slide, rels, x, y, w, h):
    n = len(rels)
    gap = 0.2
    cell_w = (w - gap * (n - 1)) / n
    for i, rel in enumerate(rels):
        cx0 = x + i * (cell_w + gap)
        path = ROOT / rel
        if not path.exists():
            placeholder(slide, cx0, y, cell_w, h, rel)
            continue
        iw, ih = Image.open(path).size
        scale = min(cell_w / iw, h / ih)
        pw, ph = iw * scale, ih * scale
        pic = slide.shapes.add_picture(str(path), Inches(cx0 + (cell_w - pw) / 2),
                                       Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
        pic.line.color.rgb = GREY
        pic.line.width = Pt(0.75)


def text_height(lines, size, w):
    est = 0.0
    for b in lines:
        per_line = max(int(w * 72 / (size * 1.05)), 1)
        est += (size / 72) * 1.5 * -(-len(b) // per_line) + 0.08
    return est


def header(slide, title):
    textbox(slide, M, 0.4, W - 2 * M, 0.8, [title], size=26, bold=True, color=DARK)
    rect(slide, M, 1.18, 1.2, 0.05, fill=ACCENT)


def content(prs, page, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=WHITE)
    header(s, page["title"])
    imgs = images_of(page)
    lead = [v for k, v in page["blocks"] if k == "lead"]
    bullets = [v for k, v in page["blocks"] if k in ("bullet", "num")]
    avail = W - 2 * M
    lead_h = text_height(lead, 22, avail) if lead else 0
    bsize = 18 if len(bullets) <= 4 else 16
    body_h = text_height(bullets, bsize, avail) if bullets else 0
    top = 1.5
    below_room = H - 0.7 - (top + lead_h + 0.2 + body_h + 0.3)
    if imgs and below_room < 2.2:
        # 左文右圖
        col = avail * 0.46
        lead_h = text_height(lead, 20, col) if lead else 0
        y = top
        if lead:
            textbox(s, M, y, col, lead_h, lead, size=20, bold=True)
            y += lead_h + 0.2
        if bullets:
            textbox(s, M, y, col, text_height(bullets, 15, col) + 0.2, bullets, size=15, bullet=True)
        place(s, imgs, M + col + 0.3, top, avail - col - 0.3, H - top - 0.75)
    else:
        y = top
        if lead:
            textbox(s, M, y, avail, lead_h, lead, size=22, bold=True)
            y += lead_h + 0.2
        if bullets:
            textbox(s, M, y, avail, body_h + 0.2, bullets, size=bsize, bullet=True)
            y += body_h + 0.3
        if imgs:
            place(s, imgs, M, y, avail, H - 0.75 - y)
    textbox(s, W - M - 1.2, H - 0.5, 1.2, 0.35, [f"{page['n']} / {total}"], size=11, color=MUTED, align="c")
    notes(s, page["notes"])


def cover(prs, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    imgs = images_of(page)
    rect(s, 0, 0, W, H, fill=IVORY)
    if imgs and (ROOT / imgs[0]).exists():
        s.shapes.add_picture(str(ROOT / imgs[0]), 0, 0, Inches(W), Inches(H))
    lines = [v for k, v in page["blocks"] if k == "bullet"]
    title = lines[0] if lines else page["title"]
    sub = lines[1:]
    textbox(s, M + 0.2, 2.5, 6.4, 1.8, [title], size=38, bold=True, color=DARK)
    if sub:
        textbox(s, M + 0.2, 4.4, 6.4, 1.0, sub, size=18, color=MUTED)
    notes(s, page["notes"])


def divider(prs, page):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, fill=IVORY)
    rect(s, M, 3.0, 1.2, 0.06, fill=ACCENT)
    textbox(s, M, 3.2, W - 2 * M, 1.2, [page["title"]], size=34, bold=True, color=DARK)
    notes(s, page["notes"])


def main():
    pages = bw4.parse(SRC)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    total = sum(1 for p in pages if not p.get("divider"))
    missing = []
    for p in pages:
        if p.get("divider"):
            divider(prs, p)
            continue
        for line in p["notes"]:
            if line.startswith("圖："):
                rel = line[3:].strip().strip("`")
                if not (ROOT / rel).exists():
                    missing.append(rel)
        if p["n"] == 1:
            cover(prs, p)
        else:
            content(prs, p, total)
    prs.save(OUT)
    print(f"{OUT.name}: {len(pages)} 頁（含 {len(pages) - total} 頁段落標題）← {SRC.relative_to(ROOT)}")
    for m in missing:
        print(f"  預留圖片（檔案不存在）：{m}")


if __name__ == "__main__":
    main()
