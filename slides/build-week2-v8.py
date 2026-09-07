#!/usr/bin/env python3
"""第二週講師版投影片，依 drafts/08-week2-slide-copy.md 第十一版（27 頁）建置。

用法：python3 slides/build-week2-v8.py  → slides/week2-agent-knowledge-work-v8.pptx
原則：標題是一句主張、畫面是證據、解釋放備註（kb/slide-design-principles-argument-to-slides.md）。
版面依 drafts/07-visual-design-rules.md；配色暖炭灰、陶土紅、鼠尾草綠；圖示 assets/week2-icons（Tabler，MIT）。
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
A = ROOT / "assets"
IC = A / "week2-icons"
OUT = ROOT / "slides" / "week2-agent-knowledge-work-v8.pptx"

DARK = RGBColor(0x2A, 0x26, 0x23)
DARK2 = RGBColor(0x3A, 0x35, 0x31)
ACCENT = RGBColor(0xB8, 0x50, 0x42)
SAGE = RGBColor(0xA7, 0xBE, 0xAE)
TINT = RGBColor(0xEE, 0xF2, 0xEF)
MUTED = RGBColor(0x6B, 0x65, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD4, 0xCE)
FONT = "PingFang TC"
W, H, M = 13.333, 7.5, 0.6

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]


# ---------- 元件 ----------
def _font(run, size, bold=False, color=DARK):
    f = run.font
    f.name, f.size, f.bold = FONT, Pt(size), bold
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def tb(s, x, y, w, h, lines, size=15, bold=False, color=DARK, align="l", anchor="t", spacing=5):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(lines, str):
        lines = [lines]
    for i, ln in enumerate(lines):
        d = ln if isinstance(ln, dict) else {"text": ln}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[d.get("align", align)]
        p.space_after = Pt(d.get("after", spacing))
        if d.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.2))))
            pPr.set("indent", str(-int(Inches(0.2))))
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        r = p.add_run()
        r.text = d["text"]
        _font(r, d.get("size", size), d.get("bold", bold), d.get("color", color))
    return box


def rect(s, x, y, w, h, fill=TINT, line=None, radius=0.08):
    shp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    return shp


def oval(s, x, y, d, fill):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    return c


def num(s, x, y, n, d=0.42, fill=ACCENT, color=WHITE, size=14):
    c = oval(s, x, y, d, fill)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    _font(r, size, True, color)
    return c


def icon(s, name, x, y, d=0.6, variant="accent", circle=None):
    """單色圖示；circle 給底色時，圖示置於圓內。"""
    if circle is not None:
        oval(s, x, y, d, circle)
        pad = d * 0.22
        s.shapes.add_picture(str(IC / f"{name}-{variant}.png"), Inches(x + pad), Inches(y + pad), Inches(d - 2 * pad), Inches(d - 2 * pad))
    else:
        s.shapes.add_picture(str(IC / f"{name}-{variant}.png"), Inches(x), Inches(y), Inches(d), Inches(d))


def picture(s, path, x, y, w, h, align="l"):
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    pw, ph = (h * ar, h) if w / h > ar else (w, w / ar)
    px = x + (w - pw) / 2 if align == "c" else x
    py = y + (h - ph) / 2 if align == "c" else y
    pic = s.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    pic.line.color.rgb = LINE
    pic.line.width = Pt(0.75)
    return px, py, pw, ph


def arrow(s, x1, y1, x2, y2, color=MUTED, dashed=False, width=1.5, head=True):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if head:
        t = etree.SubElement(ln, qn("a:tailEnd"))
        t.set("type", "triangle"); t.set("w", "med"); t.set("h", "med")
    return c


def table(s, x, y, w, rows, col_w, size=14, row_h=0.5, first_col_bold=True):
    gt = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = gt.table
    tblPr = t._tbl.tblPr
    tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.05)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            head = ri == 0
            cell.fill.fore_color.rgb = DARK if head else (WHITE if ri % 2 else TINT)
            p = cell.text_frame.paragraphs[0]
            cell.text_frame.word_wrap = True
            r = p.add_run(); r.text = str(val)
            _font(r, size, head or (first_col_bold and ci == 0), WHITE if head else DARK)
    return gt


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color


def assert_slide(title, seg, right=None):
    """白底主張頁：左上段落標，標題句 30pt 兩行內，右上可放時間。"""
    s = prs.slides.add_slide(BLANK)
    bg(s, WHITE)
    tb(s, M, 0.45, 6, 0.3, seg, size=11, color=MUTED)
    if right:
        tb(s, W - M - 4, 0.45, 4, 0.3, right, size=11, color=MUTED, align="r")
    tb(s, M, 0.85, W - 2 * M, 1.2, title, size=30, bold=True)
    return s


def card(s, x, y, w, h, head, body, ic=None, head_size=17, body_size=14, n=None):
    rect(s, x, y, w, h)
    hx = x + 0.25
    if ic:
        icon(s, ic, x + 0.25, y + 0.25, 0.6, circle=WHITE)
        hx = x + 1.0
    elif n is not None:
        num(s, x + 0.25, y + 0.3, n)
        hx = x + 0.8
    tb(s, hx, y + 0.3, w - (hx - x) - 0.2, 0.5, head, size=head_size, bold=True, anchor="m")
    tb(s, x + 0.25, y + 1.0, w - 0.5, h - 1.15, body, size=body_size)


BODY_Y = 2.25  # 主張頁內容起始


def section_slide(name, sub=None, note=""):
    s = prs.slides.add_slide(BLANK); bg(s, DARK)
    tb(s, M, 2.9, 11, 1.2, name, size=44, bold=True, color=WHITE)
    if sub:
        tb(s, W - M - 4, 6.6, 4, 0.4, sub, size=14, color=SAGE, align="r")
    notes(s, note)
    return s

def loop_crop(img, frac):
    """把直式全圖裁出某一圈的區域，回傳裁好的檔案路徑。frac 為上下比例。"""
    src = A / "week2-diagrams" / img
    dst = A / "week2-diagrams" / (img.replace(".png", f"-crop.png"))
    with Image.open(src) as im:
        h = im.height
        im.crop((0, int(h * frac[0]), im.width, int(h * frac[1]))).save(dst)
    return dst


CROPS = {"three-loops-mermaid-h1.png": (0.0, 0.318), "three-loops-mermaid-h2.png": (0.318, 0.598), "three-loops-mermaid-h3.png": (0.636, 0.915)}


def loop_slide(title, img, caption, note, seg="三個迴圈"):
    s = assert_slide(title, seg)
    if img in CROPS:
        path = loop_crop(img, CROPS[img])
        px, py, pw, ph = picture(s, path, M, BODY_Y, 7.4, 4.3, align="l")
        tx = px + pw + 0.5
    else:
        px, py, pw, ph = picture(s, A / "week2-diagrams" / img, M, BODY_Y - 0.15, 3.6, 5.5, align="l")
        tx = px + pw + 0.6
    tb(s, tx, BODY_Y + 0.4, W - M - tx, 4.5, caption, size=16, spacing=10)
    notes(s, note)
    return s


def mini_loop(s, img):
    path = loop_crop(img, CROPS[img])
    picture(s, path, W - M - 2.2, 0.4, 2.2, 1.3, align="l")

# ============ 01 標題 ============
s = prs.slides.add_slide(BLANK); bg(s, DARK)
tb(s, M, 2.3, 11, 1.2, "先玩玩看 agent 吧", size=48, bold=True, color=WHITE)
tb(s, M, 3.6, 11, 0.5, "Agent 時代的知識工作｜第二週", size=18, color=SAGE)
notes(s, "上次介紹了 agent 是什麼，這次讓大家有個方向怎麼學。前半段 30 分講怎麼合作，示範 30 分看一遍，練習 120 分自己做一份。")

# ============ 03 兩種能力 ============
s = assert_slide("AI 是放大器：領域知識 × agent 能力", "怎麼學 agent")
px, py, pw, ph = picture(s, A / "week2-evidence-pilot/magnifying-glass-book.jpg", M, BODY_Y + 0.1, 5.4, 3.4)
tb(s, M, py + ph + 0.1, 5.4, 0.35, "圖片：Julo，Wikimedia Commons，公有領域，未修改。", size=10, color=MUTED)
x = 6.6; w = W - M - x
rect(s, x, BODY_Y + 0.1, w, 1.55)
tb(s, x + 0.3, BODY_Y + 0.25, w - 0.6, 0.45, "領域知識", size=19, bold=True)
tb(s, x + 0.3, BODY_Y + 0.75, w - 0.6, 0.8, "知道方法的因果關係，知道現在缺什麼、該怎麼跟 agent 講才會得到預期的結果。", size=14)
tb(s, x, BODY_Y + 1.65, w, 0.6, "×", size=32, bold=True, color=ACCENT, align="c", anchor="m")
rect(s, x, BODY_Y + 2.3, w, 1.55)
tb(s, x + 0.3, BODY_Y + 2.45, w - 0.6, 0.45, "駕馭 agent 的能力", size=19, bold=True)
tb(s, x + 0.3, BODY_Y + 2.95, w - 0.6, 0.8, "知道不同模型擅長什麼、能做到什麼程度，偏工具的熟悉度。", size=14)
notes(s, "兩個族群。學得快沒有包袱的，缺的是領域知識；經驗老到的，不用把原本工作流程打掉換新的，分析自己工作流裡重複機械式或目標明確的部分，用 agent 試試看，也可以去補過去一直想補但沒時間補的短板。")

# ============ 04 把 AI 當同事 ============
s = assert_slide("把 AI 當同事", "怎麼學 agent")
cols = [("message-circle", "請教", ["像請教不同領域的同事：你的領域自己比較懂，但可以從它挖到很多不知道的事。", "先以自己的想法為主，不能被牽著走。"], "「我要向新進同仁介紹 GDMS，他們最先需要知道什麼？」"),
        ("list-check", "交辦", ["講明確的方向跟重要的限制：格式、範圍、長度。", "具體怎麼做不要規定太細，做法留給它。"], "「做成十頁左右可編輯的 PowerPoint，16:9，先試做一頁。」")]
cw = (W - 2 * M - 0.4) / 2
for i, (ic, h1, lines, ex) in enumerate(cols):
    x = M + i * (cw + 0.4)
    rect(s, x, BODY_Y + 0.1, cw, 3.7)
    icon(s, ic, x + 0.3, BODY_Y + 0.35, 0.7, circle=WHITE)
    tb(s, x + 1.15, BODY_Y + 0.4, cw - 1.4, 0.6, h1, size=22, bold=True, anchor="m")
    tb(s, x + 0.3, BODY_Y + 1.3, cw - 0.6, 1.5, [{"text": t, "bullet": True} for t in lines], size=15, spacing=8)
    tb(s, x + 0.3, BODY_Y + 2.7, cw - 0.6, 1.0, ex, size=14, bold=True, color=ACCENT)
tb(s, M, BODY_Y + 4.05, W - 2 * M, 0.4, "最省事的交辦：只給一句話、不給材料、不請教。示範開頭會先這樣做一次。", size=14, color=MUTED)
notes(s, "太籠統會沒跟自己對齊，太細節可能會錯過更好的方法。已有自己累積的工作流程，不用急著全部切換。最省事的交辦是只給一句話，示範開頭會先這樣做一次。")

# ============ 05 從自己需要補足的地方開始 ============
s = assert_slide("從自己需要補足的地方開始", "怎麼學 agent")
helps = [("compass", "缺知識與方向", "請它說明做法，再自行查閱。"), ("tool", "有想法但不熟工具", "描述想法請它實作，再試用修改，像 vibe coding。"),
         ("settings", "簡單重複的雜工", "交代規則與完成條件，交給它處理。"), ("zoom-check", "需要人抓盲點", "當作審查者，讓它找遺漏，再自行決定是否修改。")]
cw = (W - 2 * M - 0.4) / 2
for i, (ic, h1, body) in enumerate(helps):
    x = M + (i % 2) * (cw + 0.4); y = BODY_Y + 0.1 + (i // 2) * 2.2
    card(s, x, y, cw, 2.0, h1, body, ic=ic, head_size=19, body_size=15)
notes(s, "同一件事可能同時需要幾種協助。")

# ============ 06 交付責任 ============
s = assert_slide("交付出去的東西是要負責的", "怎麼學 agent")
lines = ["用幾句話完成什麼東西看起來很厲害，但那不是你要的東西。", "與其 AI 一次吐一堆看不完的東西，", "不如一點一點的把工作流程替換。"]
for i, t in enumerate(lines):
    tb(s, M + 0.2, BODY_Y + 0.3 + i * 1.15, 6.4, 1.0, t, size=20, bold=(i == 2), anchor="m")
px, py, pw, ph = picture(s, A / "week2-baseline/baseline-method-card.png", 7.4, BODY_Y + 0.2, W - M - 7.4, 3.0)
tb(s, 7.4, py + ph + 0.1, W - M - 7.4, 0.7, "一句話生成的簡報第 4 頁：版面漂亮、三塊說明塞滿，講的人卻答不出第四種方法是什麼。", size=12, color=MUTED)
notes(s, "對照組的反應：做得不錯，但拿去報就會卡，因為不知道第四種方法是什麼，被問就答不出來。漂亮這點可以吸收，後面第三圈會用到。")

# ============ 07 外部化：資料夾 ============
s = assert_slide("東西都要外部化存起來", "怎麼學 agent")
tree = [("AGENTS.md", "工作方式與判斷標準，每次對話都讀", False), ("工作紀錄.md", "討論脈絡與每次定案", False),
        ("kb/", "找到的材料與洞見；arguments 論點類、tools 工具類", True), ("drafts/", "教案、逐頁稿、棄案，可以大改", True),
        ("skills/", "定案後的判斷標準，一個主題一檔", True), ("assets/", "圖片與圖示，附來源", True),
        ("slides/", "投影片與建置腳本", True), ("lessons/", "學員看得到的文章版教材", True), ("tmp/", "試跑與暫存，不進 git", True)]
x = M + 0.4
for i, (name, desc, folder) in enumerate(tree):
    y = BODY_Y - 0.1 + i * 0.5
    if folder:
        icon(s, "folder", x, y + 0.05, 0.36)
    else:
        rect(s, x + 0.06, y + 0.08, 0.26, 0.32, fill=WHITE, line=MUTED, radius=0.05)
    tb(s, x + 0.55, y, 2.6, 0.46, name, size=15, bold=True, anchor="m")
    tb(s, x + 3.2, y, W - M - x - 3.2, 0.46, desc, size=13, color=MUTED, anchor="m")
tb(s, M, 6.95, W - 2 * M, 0.35, "這門課 repo 的現況；結構是跟 agent 互動長出來的，kb 昨天才分成兩類。", size=11, color=MUTED)
notes(s, "整個資料夾結構是有機長出來的，只要給 agent 這個概念：東西都外部化存起來。kb 是常見的處理方式，其他依專案需求跟著互動建立；有自己熟悉的整理方式，叫 agent 先照抄也可以。")

# ============ 08 段：三個迴圈 ============
section_slide("三個迴圈", note="內容不是一次寫成的。先一張全圖，之後每圈兩頁：一頁講那一圈怎麼走，一頁用這份投影片當實例。")

# ============ 09 全圖 ============
loop_slide("內容要走三個迴圈", "three-loops-mermaid.png",
           [{"text": "確認方向、建立架構、調整風格，三圈各有人的判斷（紅框）。", "after": 12},
            {"text": "灰色虛線在圈內回頭；紅色虛線退回上一圈。", "after": 12},
            {"text": "跟軟體開發對得上：需求探索、建立架構、寫程式。", "color": MUTED}],
           "前半段是訂方向，後半段是建立架構，會看到兩個大的迴圈；內容確定了還有第三圈，讓交付物對齊自己的標準。接下來每圈兩頁。")

# ============ 10 確認方向 ============
loop_slide("確認方向", "three-loops-mermaid-h1.png",
           [{"text": "人看過所有探索的資料，先不用 AI 寫一版大綱。", "bold": True, "after": 12},
            {"text": "一方面人先全部想過一遍，另一方面 AI 會順著這個方向寫，減少認知債。", "after": 12},
            {"text": "真的想不到，再叫 AI 串串看。", "color": MUTED}],
           "有自己的想法就先以自己的為主，洞見從這裡插進去，成果才代表自己。後面才看到的東西可能推翻論點，就回去調整方向，但會隨著交付日期接近而鎖定。")

# ============ 11 實例：方向 ============
s = assert_slide("確認方向：方向改了三次才定下來", "三個迴圈")
dates = [("9 月 2 日", "文件專案化與模型能力"), ("9 月 4 日", "公部門用 AI 的阻礙與規範"), ("9 月 6 日", "先體驗 agent")]
for i, (d, t) in enumerate(dates):
    y = BODY_Y + 0.15 + i * 1.45
    rect(s, M, y, 5.6, 1.1, fill=(WHITE if i < 2 else TINT), line=(LINE if i < 2 else None))
    tb(s, M + 0.3, y + 0.15, 1.6, 0.8, d, size=13, color=MUTED, anchor="m")
    tb(s, M + 1.9, y + 0.15, 3.5, 0.8, t, size=17, bold=(i == 2), anchor="m")
    if i < 2:
        arrow(s, M + 2.8, y + 1.1, M + 2.8, y + 1.45, color=ACCENT, width=2)
tb(s, M, BODY_Y + 4.5, 5.6, 0.4, "確認方向那一圈：方向定了，材料進來又改細節", size=12, color=MUTED)
x = 7.0; w = W - M - x
rect(s, x, BODY_Y + 0.15, w, 4.3)
tb(s, x + 0.3, BODY_Y + 0.35, w - 0.6, 0.4, "材料進來又改的", size=14, bold=True, color=ACCENT)
moved = [("客服 AI 研究", "找來又移出，不好講"), ("簡報形式比較與 NESA", "從題目退成起點，再整段移出"), ("臺大兩個教學案例", "退到延伸閱讀"), ("業務簡報一句話對照組", "移到示範開頭，換成 GDMS")]
for i, (a, b) in enumerate(moved):
    y = BODY_Y + 0.95 + i * 0.85
    tb(s, x + 0.3, y, w - 0.6, 0.35, a, size=15, bold=True)
    tb(s, x + 0.3, y + 0.36, w - 0.6, 0.4, b, size=13, color=MUTED)
notes(s, "這是確認方向那一圈。連兩個迴圈的圖本身，也是從「簡報內容逐漸成熟的流程」改了五次才定型。")

# ============ 12 建立架構 ============
loop_slide("建立架構", "three-loops-mermaid-h2.png",
           [{"text": "把 AI 說的話順成自己說的話，把真實的肉長出來。", "bold": True, "after": 12},
            {"text": "迭代回圈小一點，內容盡量留在 Markdown；甚至先在對話裡叫 agent 講一遍怎麼做，確定了再落檔。", "after": 12},
            {"text": "講稿像文章，投影片分口述跟畫面，節奏不一樣。", "color": MUTED}],
           "一開始可以先叫 AI 打個草稿，這裡需要人大量參與，補文獻來源或圖片。越快進到具體改起來越慢。")

# ============ 12 實例：補進去的想法 ============
s = assert_slide("建立架構：講稿裡補進去的想法", "三個迴圈")
ideas = [("領域知識 × agent 能力，兩邊都要有", "9 月 6 日"), ("兩個族群，各缺一邊", "9 月 6 日"), ("聊天像請教導師，交辦要講明限制", "9 月 6 日"),
         ("人先不用 AI 寫大綱，減少認知債", "9 月 6 日"), ("迭代回圈小一點，內容留在 md", "9 月 6 日"), ("產出物是 kb、draft、ppt", "9 月 7 日"),
         ("第三個迴圈，讓交付物對齊自己的標準", "9 月 8 日")]
for i, (t, d) in enumerate(ideas):
    y = BODY_Y + 0.05 + i * 0.62
    num(s, M, y + 0.07, i + 1, d=0.4, size=13)
    tb(s, M + 0.6, y, 6.6, 0.55, t, size=16, anchor="m")
    tb(s, M + 7.3, y, 1.4, 0.55, d, size=12, color=MUTED, anchor="m")
rect(s, 9.2, BODY_Y + 0.6, W - M - 9.2, 2.6, fill=TINT)
tb(s, 9.5, BODY_Y + 0.85, W - M - 9.8, 2.2, ["這些段落 agent 原本沒有，", "是討論時補進講稿的。"], size=17, bold=True, spacing=6)
notes(s, "我中間一直在來回修那個講稿。agent 沿論點展開的是通順的字，觀點是人插進去的。改語氣是另一回事，在第六步。")

# ============ 14 調整風格 ============
loop_slide("調整風格", "three-loops-mermaid-h3.png",
           [{"text": "內容確定了，讓交付物對齊自己的標準。", "bold": True, "after": 12},
            {"text": "一直在改，就是要去找一個原則：叫 agent 先整理別人的做法照著做，這就是補短板。", "after": 12},
            {"text": "依序：分頁與順序 → 文字與語氣 → 版面與圖文。", "color": ACCENT, "bold": True}],
           "入口是標準。自己有判斷力的面向直接拿規則檔，用過有效的補進 AGENTS.md，這一格就變短。AI 語氣也是這一圈的事。")

# ============ 14 實例：分頁動作 ============
s = assert_slide("調整風格：分頁時實際做了什麼", "三個迴圈")
acts = [("換成原話當底", "生成的「還沒有方法時請教」，改成「聊天像請教導師」"), ("只讀標題連讀，刪頁", "研究專案情境頁、時間表頁"),
        ("換順序", "把 AI 當同事前移；外部化接交付責任；迴圈與實例穿插"), ("縮標題", "「開空專案，建 kb 與 AGENTS.md」變「開空專案」"),
        ("加段落頁", "三個迴圈、示範、練習"), ("再看逐頁稿", "每頁補一行目的，畫面照它選")]
for i, (h1, ex) in enumerate(acts):
    y = BODY_Y + 0.05 + i * 0.72
    num(s, M, y + 0.1, i + 1, d=0.4, size=13)
    tb(s, M + 0.6, y, 3.2, 0.6, h1, size=16, bold=True, anchor="m")
    tb(s, M + 3.9, y, W - M - (M + 3.9), 0.6, ex, size=14, color=MUTED, anchor="m")
tb(s, M, 6.7, W - 2 * M, 0.4, "現在正在做的就是這一圈：畫面是最後一步，內容問題都在前兩步解決。", size=13, color=ACCENT)
notes(s, "先把 ghost deck 寫好，用原話來調整 ghost deck 比較能看懂；標題不要有逗號的長度；這頁其實不用；五跟四反過來看看。")

# ============ 16 段：示範 ============
section_slide("示範：向新進同仁介紹 GDMS", sub="示範 30 分", note="講師走一遍八步，只展開一條論點。學員記兩件事：自己會加哪個條件，最想先試哪一步。")

# ============ 17 開始前 ============
s = assert_slide("開始前先確認資料去哪裡", "示範與練習")
pre = [("shield", "資料使用設定", "關閉訓練不代表資料不上雲。"), ("lock", "repo 公開與否", "private repo 同樣在雲端。"),
       ("route", "PPTX 製作路徑", "這台機器用什麼做出可編輯的檔案，講師告知。"), ("eye", "預覽方式", "無預覽軟體時，請 agent 匯出 PDF。")]
cw = (W - 2 * M - 0.4) / 2
for i, (ic, h1, body) in enumerate(pre):
    x = M + (i % 2) * (cw + 0.4); y = BODY_Y + 0.1 + (i // 2) * 2.0
    card(s, x, y, cw, 1.8, h1, body, ic=ic, head_size=18, body_size=14)
tb(s, M, 6.4, W - 2 * M, 0.5, "同時把「幫我做一份向新進同仁介紹 GDMS 的簡報」貼給 agent，讓它在背景做第一版。接下來八頁，示範時看，練習時照做。", size=14, color=ACCENT, bold=True)
notes(s, "要把提供訓練關掉，push GitHub 本身就是上雲，測試的資料不要選機密的。示範只展開一條論點。接下來八頁，示範時看，練習時照做。")

# ============ 18–25 八步：固定版式 ============
def step_slide(n, title, do, say, check, demo, prac, evidence, note):
    s = assert_slide(f"第{n}步：{title}", "示範與練習", right=f"示範 {demo} 分　練習 {prac} 分")
    lx, lw = M, 6.9
    blocks = [("做什麼", do), ("對 agent 說", say), ("確認什麼", check)]
    y = BODY_Y
    hs = [1.55, 1.25, 1.35]
    for (lab, lines), h in zip(blocks, hs):
        rect(s, lx, y, lw, h)
        tb(s, lx + 0.25, y + 0.15, 1.4, 0.35, lab, size=12, bold=True, color=ACCENT)
        if lab == "對 agent 說":
            tb(s, lx + 0.25, y + 0.5, lw - 0.5, h - 0.6, lines, size=14, bold=True, color=DARK, spacing=4)
        else:
            tb(s, lx + 0.25, y + 0.5, lw - 0.5, h - 0.6, [{"text": t, "bullet": True} for t in lines], size=14, spacing=4)
        y += h + 0.15
    rx = lx + lw + 0.4
    evidence(s, rx, BODY_Y, W - M - rx, y - 0.15 - BODY_Y)
    notes(s, note)
    return s

def ev14(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "練習專案／", size=14, bold=True)
    tb(s, x + 0.3, y + 0.7, w - 0.6, 2.5, [{"text": "AGENTS.md", "bold": True, "after": 2}, {"text": "保存方式與共用要求", "size": 12, "color": MUTED, "after": 14},
                                          {"text": "kb/第一次討論.md", "bold": True, "after": 2}, {"text": "題目、受眾、用途、想法、未定事項", "size": 12, "color": MUTED}], size=15)
    icon(s, "folder", x + w - 1.3, y + h - 1.3, 1.0)

def ev16(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "方向修正的例子", size=14, bold=True, color=ACCENT)
    tb(s, x + 0.3, y + 0.7, w - 0.6, 0.9, "「圖片好看但不能改」", size=15, color=MUTED)
    tb(s, x + 0.3, y + 1.35, w - 0.6, 0.4, "↓ 看過故事線", size=12, color=MUTED)
    tb(s, x + 0.3, y + 1.75, w - 0.6, 1.0, "「一次性、不再修改的內容才選圖片」", size=15, bold=True)
    tb(s, x + 0.3, y + 2.9, w - 0.6, 1.0, "棄用「互動性」比較：受眾不需要，也未實測。", size=13, color=MUTED)

def ev17(s, x, y, w, h):
    rows = [["修改前", "HTML 改完可直接存檔，同事能繼續改（試跑例）"], ["對照來源", "預設只存瀏覽器本機，寫回需另備 server"],
            ["修改後", "可在瀏覽器改文字與位置；預設只存該台電腦，交接須先匯出。作者說明，未實測"]]
    table(s, x, y, w, rows, [1.2, w - 1.2], size=12, row_h=1.1)

def ev19(s, x, y, w, h):
    ph = (h - 0.9) / 2
    tb(s, x, y, w, 0.3, "首次輸出：整頁太空、字小", size=12, color=MUTED)
    picture(s, A / "week2-student-run/first-output-page1.png", x, y + 0.35, w, ph)
    tb(s, x, y + ph + 0.55, w, 0.3, "修改後：標題加大、關鍵字做成並排色塊", size=12, color=ACCENT, bold=True)
    picture(s, A / "week2-student-run/revised-page1.png", x, y + ph + 0.9, w, ph)

def ev20(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "agent 無法代做", size=14, bold=True, color=ACCENT)
    items = [("pointer", "親自點選文字，確認可編輯"), ("microphone", "實際講一頁，計時"), ("users", "請人重述重點")]
    for i, (ic, t) in enumerate(items):
        yy = y + 0.85 + i * 1.05
        icon(s, ic, x + 0.3, yy, 0.7, circle=TINT)
        tb(s, x + 1.2, yy, w - 1.5, 0.7, t, size=15, anchor="m")

def ev14(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "練習專案／", size=14, bold=True)
    tb(s, x + 0.3, y + 0.7, w - 0.6, 2.5, [{"text": "AGENTS.md", "bold": True, "after": 2}, {"text": "保存方式與共用要求", "size": 12, "color": MUTED, "after": 14},
                                          {"text": "kb/第一次討論.md", "bold": True, "after": 2}, {"text": "題目、受眾、用途、想法、未定事項", "size": 12, "color": MUTED}], size=15)
    icon(s, "folder", x + w - 1.3, y + h - 1.3, 1.0)

def ev16(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "方向修正的例子", size=14, bold=True, color=ACCENT)
    tb(s, x + 0.3, y + 0.7, w - 0.6, 0.9, "「圖片好看但不能改」", size=15, color=MUTED)
    tb(s, x + 0.3, y + 1.35, w - 0.6, 0.4, "↓ 看過故事線", size=12, color=MUTED)
    tb(s, x + 0.3, y + 1.75, w - 0.6, 1.0, "「一次性、不再修改的內容才選圖片」", size=15, bold=True)
    tb(s, x + 0.3, y + 2.9, w - 0.6, 1.0, "棄用「互動性」比較：受眾不需要，也未實測。", size=13, color=MUTED)

def ev17(s, x, y, w, h):
    rows = [["修改前", "HTML 改完可直接存檔，同事能繼續改（試跑例）"], ["對照來源", "預設只存瀏覽器本機，寫回需另備 server"],
            ["修改後", "可在瀏覽器改文字與位置；預設只存該台電腦，交接須先匯出。作者說明，未實測"]]
    table(s, x, y, w, rows, [1.2, w - 1.2], size=12, row_h=1.1)

def ev19(s, x, y, w, h):
    ph = (h - 0.9) / 2
    tb(s, x, y, w, 0.3, "首次輸出：整頁太空、字小", size=12, color=MUTED)
    picture(s, A / "week2-student-run/first-output-page1.png", x, y + 0.35, w, ph)
    tb(s, x, y + ph + 0.55, w, 0.3, "修改後：標題加大、關鍵字做成並排色塊", size=12, color=ACCENT, bold=True)
    picture(s, A / "week2-student-run/revised-page1.png", x, y + ph + 0.9, w, ph)

def ev20(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "agent 無法代做", size=14, bold=True, color=ACCENT)
    items = [("pointer", "親自點選文字，確認可編輯"), ("microphone", "自己找時間真的講一遍，計時")]
    for i, (ic, t) in enumerate(items):
        yy = y + 0.85 + i * 1.05
        icon(s, ic, x + 0.3, yy, 0.7, circle=TINT)
        tb(s, x + 1.2, yy, w - 1.5, 0.7, t, size=15, anchor="m")

def ev_gdms(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=ACCENT)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "預設題的材料", size=14, bold=True, color=ACCENT)
    tb(s, x + 0.3, y + 0.7, w - 0.6, h - 0.9, [{"text": "只用講師提供的 GDMS 公開頁面與說明文件。", "bullet": True},
                                              {"text": "需要帳號才看得到的畫面與資料，不提供給工具。", "bullet": True},
                                              {"text": "agent 寫的功能、資料範圍與數字，一律對照說明頁。", "bullet": True}], size=14, spacing=8)

def ev_rules(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "語氣的判準來源", size=14, bold=True, color=ACCENT)
    for i, t in enumerate(["能用自己說過的原話就用原話", "自己的文案語氣規則", "維基百科 AI 寫作特徵清單", "OpenAI 禁用詞、Anthropic 的 mannered prose"]):
        yy = y + 0.8 + i * 0.7
        num(s, x + 0.3, yy + 0.03, i + 1, d=0.38, size=12)
        tb(s, x + 0.85, yy, w - 1.1, 0.45, t, size=14, anchor="m")
    tb(s, x + 0.3, y + h - 0.7, w - 0.6, 0.5, "請 agent 標出沒有原話依據的句子。", size=12, color=MUTED)

def ev_ghost(s, x, y, w, h):
    rect(s, x, y, w, h, fill=WHITE, line=LINE)
    tb(s, x + 0.3, y + 0.25, w - 0.6, 0.35, "這份投影片的 ghost deck 前幾行", size=14, bold=True, color=ACCENT)
    lines = ["2 研究專案裡的材料、想法與成果持續累積", "3 你卡在哪裡，由兩種能力決定", "4 缺的是哪種能力，就需要哪種協助", "5 還沒有方法時請教，已知道要做什麼時交辦", "6 先確認方向，再建立架構，最後排版", "7 過程保存成三類檔案"]
    tb(s, x + 0.3, y + 0.75, w - 0.6, h - 1.0, [{"text": t, "after": 6} for t in lines], size=12)

num_cn = ["一", "二", "三", "四", "五", "六", "七", "八"]

step_slide(num_cn[0], "開空專案",
           ["講這個專案大致上要做什麼。", "建立 kb，叫 agent 把討論都寫進 kb，這個動作寫進 AGENTS.md。"],
           "「我要向新進同仁介紹 GDMS。先建 kb/ 保存材料與討論，保存方式寫進 AGENTS.md。」",
           ["開啟兩個檔案，太長就刪到只留保存位置。", "再請它讀取 AGENTS.md。"], 2, 5, ev14,
           "有自己熟悉的整理方式就叫 agent 照抄。開啟確認、刪短、明確請它讀取，這三條來自試跑。")

step_slide(num_cn[1], "講背景與限制",
           ["講受眾、時間、簡報長度這些限制條件。", "拓展主題看要讓觀眾知道什麼。", "找到的材料整理進 kb，附來源與日期。"],
           "「受眾是還沒用過 GDMS 的新進同仁。先讀我給的公開頁面，每一份整理進 kb。」",
           ["每份實際開啟；區分作者說明、自己看過、待確認。", "30 分停止搜尋。"], 3, 20, ev_gdms,
           "用聊天的方式把背景講一講，問問看 agent 建議，最後落成 md 存起來。任何一步卡住都可以再找一輪材料。")

step_slide(num_cn[2], "先找出想講的方向",
           ["先想清楚這次想講的方向和理由。", "有了大方向叫 AI 拓展看看有沒有東西要調整，有就繼續聊補齊。", "棄用的另存一份，記理由。"],
           "「我覺得新進同仁最先要知道的是＿＿，因為＿＿。照這個順序整理故事線。」",
           ["連讀故事線，至少修正一處。"], 3, 15, ev16,
           "先不用 AI 寫一個大綱出來；真的想不到，叫 AI 串串看也可以。右側例子取自試跑，題目為簡報形式比較。")

step_slide(num_cn[3], "起草講稿",
           ["請 agent 先寫一段，看過再展開。", "把 AI 說的話順成自己說的話，補文獻來源或圖片。", "挑一句斷言請它對照說明頁，前後版本都留。"],
           "「先寫一段給我看。」　「這句『＿＿』請對照說明頁檢查，前後版本都留在草稿裡。」",
           ["空泛結尾要補條件。", "未實測照標未實測。"], 4, 15, ev17,
           "講稿比較像文章，可以作為輔助使用。挑一句查證是試跑後加的動作，學員不會主動做，第 50 分鐘全場再示範一次。")

step_slide(num_cn[4], "找分頁原則",
           ["請 agent 找投影片該怎麼分頁的原則，挑出要用的。", "先只列每頁一句標題，連讀通過才往下。", "再逐頁配證據與口述，存成分頁文案。"],
           "「先只列每頁一句標題，我連讀過再配內容。」",
           ["只讀標題能講完故事。", "每頁對應草稿一段，對不上的刪；第一頁不能省。"], 2, 10, ev_ghost,
           "分頁方法找到好幾種，本課選綜合的一種。先讓 agent 自己找一輪再對照課程筆記。講稿更具體後做成分頁草稿，投影片分口述跟畫面。")

step_slide(num_cn[5], "修改語氣",
           ["能用自己說過的原話就用原話。", "其餘依文案語氣規則改，沒有規則就先請 agent 整理 AI 寫作特徵清單。", "自己讀一遍，不像自己講話的句子指出來再改。"],
           "「照這份規則逐頁改寫，標出哪些句子沒有原話依據。」",
           ["還在 Markdown 上改。", "看畫面最刺眼的口號句在這一步解決。"], 2, 10, ev_rules,
           "文字語氣是第三圈的第二個面向，便宜，先做。不是精簡成奇怪的金句，要是書面文。")

step_slide(num_cn[6], "找排版原則",
           ["拿出視覺設計規則，或請 agent 再找一輪比對。", "第一次先做一頁確認做得出來；之後整份產出，匯出 PDF 整份看。", "版面問題改規則重建；內容問題退回第五、六步。"],
           "「照這份分頁文案做成可編輯的 PowerPoint，16:9，版面照這份規則，做完匯出 PDF 給我看。」",
           ["備註有口述與來源；文字都是文字框。"], 7, 25, ev19,
           "排版就借用一下 ppt 的技巧。迭代回圈小一點，才不會修太多美編排版，方向改變整個丟掉返工。漂亮這點可以吸收。")

step_slide(num_cn[7], "整個過一遍",
           ["自己找時間真的講一遍，計時。", "講不順的那頁，回草稿補或刪。"],
           "「估每頁口述的字數和時間。」",
           ["開頭說的需求，每項都在成品裡。", "kb/、草稿、分頁文案、PPTX 最新版找得到。"], 5, 10, ev20,
           "最後人整個過一遍練習一下去微調，算是一個確認，不用兩人一組。")

# ============ 26 段：練習 ============
section_slide("練習", sub="練習 120 分", note="自由發揮，不收成品。")

# ============ 27 練習 ============
s = assert_slide("練習：題目自選", "練習 120 分")
cols = [("tag", "題目", ["預設題可用。", "自選題：有受眾與用途、材料公開或可虛構、兩小時內找得到依據。"]),
        ("player-play", "開始時", ["題目一句話交給 agent 做第一版，背景執行。", "另開視窗走八步。"]),
        ("stack-2", "成品", ["10 到 12 頁可編輯 PPTX。", "每頁對應草稿一段。"])]
cw = (W - 2 * M - 0.6) / 3
for i, (ic, h1, lines) in enumerate(cols):
    x = M + i * (cw + 0.3)
    rect(s, x, BODY_Y + 0.1, cw, 4.0)
    icon(s, ic, x + 0.3, BODY_Y + 0.35, 0.7, circle=WHITE)
    tb(s, x + 1.15, BODY_Y + 0.4, cw - 1.4, 0.6, h1, size=21, bold=True, anchor="m")
    tb(s, x + 0.3, BODY_Y + 1.3, cw - 0.6, 2.7, [{"text": t, "bullet": True} for t in lines], size=14, spacing=8)
notes(s, "第一版留到回看。第 10 分鐘講師巡一輪，寫不出受眾與用途者改用預設題。")

# ============ 28 回看 ============
s = assert_slide("回看：第一版和自己做的並排", "練習 120 分")
for i, (lab, sub) in enumerate([("第一版", "一句話生成"), ("自己做的", "八步之後")]):
    x = M + 1.0 + i * 6.4
    rect(s, x, BODY_Y + 0.1, 4.6, 2.2)
    icon(s, "stack-2", x + 0.4, BODY_Y + 0.55, 1.3)
    tb(s, x + 2.0, BODY_Y + 0.6, 2.5, 0.6, lab, size=24, bold=True, anchor="m")
    tb(s, x + 2.0, BODY_Y + 1.25, 2.5, 0.5, sub, size=14, color=MUTED)
icon(s, "arrows-left-right", W / 2 - 0.45, BODY_Y + 0.75, 0.9, circle=TINT)
qs = [("補上了什麼", "查到的材料、自己的例子、修正過的斷言"), ("哪些是自己的判斷", "採用或刪除的論點、棄案理由、開頭在意的條件"), ("哪些交給了 agent", "搜尋、整理、查證、排版")]
for i, (h1, body) in enumerate(qs):
    y = BODY_Y + 2.7 + i * 0.75
    num(s, M + 1.0, y + 0.05, i + 1, d=0.4, size=13)
    tb(s, M + 1.6, y, 3.6, 0.5, h1, size=16, bold=True, anchor="m")
    tb(s, M + 5.3, y, 7, 0.5, body, size=14, color=MUTED, anchor="m")
notes(s, "說不出理由的那頁，內容還不是自己的。核心訊息未定，目前最接近的一句：kb 會自己長出來，收斂想法很快，因為 agent 會自己搜尋；後面 ppt 製作更明顯，agent 會自己把簡報完成。")

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
