#!/usr/bin/env python3
"""第四週投影片建置：把三份逐頁稿（drafts/09、10、11）各產成一份素版 PPTX。

用法：python3 slides/build-week4.py            → slides/week4-1.pptx、week4-2.pptx、week4-3.pptx
      python3 slides/build-week4.py 09         → 只建其中一份

逐頁稿格式（drafts/09-week4-1-slides.md 等）：
- 「## 逐頁」之後，每頁以「### N｜標題」開頭；頁內可用「**畫面**」「**筆記**」「**原話與來源**」「**圖規格**」分區，畫面區上投影片，其餘進備註。舊格式（無分區，以「- 原話：」等前綴區分）仍支援。
- 內文：一般條列「- 」、編號「1. 」、Markdown 表格、```text 程式區塊，依序放到畫面。
- 「- 原話：」「- 口述：」「- 【無原話】」「- 案例：」「- 圖：」「- 來源：」不上畫面，進備註。
- 「- 引文頁內文：『…』」以引文方塊放到畫面。
- 第 1 頁視為封面。

素版原則：白底、標題一句、內文條列或表格、備註放原話與口述；不做視覺設計，供使用者先有東西改。
"""
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DECKS = {
    "09": ("drafts/09-week4-1-slides.md", "week4-1.pptx"),
    "10": ("drafts/10-week4-2-slides.md", "week4-2.pptx"),
    "11": ("drafts/11-week4-3-slides.md", "week4-3.pptx"),
}

DARK = RGBColor(0x2A, 0x26, 0x23)
MUTED = RGBColor(0x6B, 0x65, 0x60)
ACCENT = RGBColor(0xB8, 0x50, 0x42)
TINT = RGBColor(0xEE, 0xF2, 0xEF)
LINE = RGBColor(0xD9, 0xD4, 0xCE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "PingFang TC"
W, H, M = 13.333, 7.5, 0.7

NOTE_PREFIXES = ("原話：", "口述：", "【無原話】", "案例：", "備註用：", "圖：", "來源：", "圖格：")


# ---------- 解析 ----------
def parse(md_path):
    text = Path(md_path).read_text(encoding="utf-8")
    text = text[text.index("## 逐頁"):]
    pages = []
    cur = None
    in_code = False
    section = "畫面"
    for line in text.split("\n"):
        m = re.match(r"^### (\d+)｜(.+)$", line)
        if m:
            cur = {"n": int(m.group(1)), "title": m.group(2).strip(),
                   "blocks": [], "notes": []}
            pages.append(cur)
            section = "畫面"
            continue
        if cur is None:
            continue
        m = re.match(r"^\*\*(畫面|筆記|原話與來源|圖規格)\*\*$", line.strip())
        if m:
            section = m.group(1)
            if section != "畫面":
                cur["notes"].append(f"【{section}】")
            continue
        if section != "畫面":
            t = line.strip()
            if t and t != "無。":
                cur["notes"].append(re.sub(r"^- ", "", t))
            continue
        if line.startswith("```"):
            if in_code:
                in_code = False
            else:
                in_code = True
                cur["blocks"].append(("code", []))
            continue
        if in_code:
            cur["blocks"][-1][1].append(line)
            continue
        s = line.strip()
        if not s:
            continue
        if s.startswith("|"):
            if s.startswith("|---") or set(s.replace("|", "").strip()) <= set("-: "):
                continue
            cells = [c.strip() for c in s.strip("|").split("|")]
            if cur["blocks"] and cur["blocks"][-1][0] == "table":
                cur["blocks"][-1][1].append(cells)
            else:
                cur["blocks"].append(("table", [cells]))
            continue
        m = re.match(r"^- (.*)$", s)
        if m:
            body = m.group(1).strip()
            if body.startswith("引文頁內文："):
                q = body[len("引文頁內文："):].strip()
                q = re.sub(r"（[^）]*）$", "", q).strip().strip("「」")
                cur["blocks"].append(("quote", q))
                continue
            if body.startswith(NOTE_PREFIXES):
                cur["notes"].append(body)
                continue
            cur["blocks"].append(("bullet", body))
            continue
        m = re.match(r"^(\d+)\. (.*)$", s)
        if m:
            cur["blocks"].append(("num", f"{m.group(1)}. {m.group(2).strip()}"))
            continue
        cur["blocks"].append(("bullet", s))
    return pages


# ---------- 元件 ----------
def _run(run, size, bold=False, color=DARK):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def textbox(slide, x, y, w, h, lines, size=18, bold=False, color=DARK,
            align="l", anchor="t", spacing=6, bullet=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER}[align]
        p.space_after = Pt(spacing)
        txt = ("•  " + ln) if bullet and not re.match(r"^\d+\. ", ln) else ln
        r = p.add_run()
        r.text = txt
        _run(r, size, bold, color)
    return box


def rect(slide, x, y, w, h, fill=TINT, line=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    return shp


def table(slide, x, y, w, rows, size=13, row_h=0.42):
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    # 欄寬依內容長度分配
    lens = [max(len(r[c]) for r in rows) for c in range(ncol)]
    lens = [max(l, 4) for l in lens]
    total = sum(lens)
    col_w = [w * l / total for l in lens]
    # 估行高：每行字數超過欄寬可容納量就加高
    heights = []
    for r in rows:
        mx = 1
        for c, cell in enumerate(r):
            per_line = max(int(col_w[c] * 72 / (size * 1.05)), 1)
            mx = max(mx, -(-len(cell) // per_line))
        heights.append(row_h * (0.55 + 0.45 * mx))
    shp = slide.shapes.add_table(len(rows), ncol, Inches(x), Inches(y), Inches(w), Inches(sum(heights)))
    tbl = shp.table
    for c in range(ncol):
        tbl.columns[c].width = Inches(col_w[c])
    for i, r in enumerate(rows):
        tbl.rows[i].height = Inches(heights[i])
        for c, cell in enumerate(r):
            tc = tbl.cell(i, c)
            tc.margin_left = tc.margin_right = Inches(0.06)
            tc.margin_top = tc.margin_bottom = Inches(0.03)
            tc.fill.solid()
            tc.fill.fore_color.rgb = TINT if i == 0 else WHITE
            tf = tc.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = cell.replace("**", "")
            _run(run, size, bold=(i == 0 or c == 0), color=DARK)
    return sum(heights)


def notes(slide, lines):
    if not lines:
        return
    slide.notes_slide.notes_text_frame.text = "\n".join(lines)


# ---------- 版面 ----------
def cover(prs, page, deck_label):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    lines = [b[1] for b in page["blocks"] if b[0] == "bullet"]
    title = lines[0] if lines else page["title"]
    sub = lines[1:] if len(lines) > 1 else [deck_label]
    rect(s, 0, 0, 0.35, H, fill=ACCENT)
    textbox(s, M + 0.4, 2.3, W - 2 * M - 0.4, 1.6, [title], size=36, bold=True)
    textbox(s, M + 0.4, 4.0, W - 2 * M - 0.4, 1.2, sub, size=18, color=MUTED)
    notes(s, page["notes"])
    return s


def content(prs, page, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(s, M, 0.45, W - 2 * M, 0.9, [page["title"]], size=28, bold=True)
    ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(M), Inches(1.32), Inches(W - 2 * M), Inches(0.03))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT; ln.line.fill.background()
    y = 1.6
    avail_w = W - 2 * M
    bullets = []

    def flush_bullets():
        nonlocal y, bullets
        if not bullets:
            return
        n = len(bullets)
        size = 18 if n <= 5 else 16 if n <= 7 else 14
        est = 0.0
        for b in bullets:
            per_line = max(int(avail_w * 72 / (size * 1.05)), 1)
            est += (size / 72) * 1.5 * -(-len(b) // per_line) + 0.08
        textbox(s, M, y, avail_w, est + 0.2, bullets, size=size, bullet=True)
        y += est + 0.25
        bullets = []

    for kind, val in page["blocks"]:
        if kind in ("bullet", "num"):
            bullets.append(val)
            continue
        flush_bullets()
        if kind == "table":
            nrow = len(val)
            size = 13 if nrow <= 7 else 12 if nrow <= 9 else 11
            h = table(s, M, y, avail_w, val, size=size)
            y += h + 0.25
        elif kind == "code":
            lines = [l for l in val]
            h = 0.32 * len(lines) + 0.3
            rect(s, M, y, avail_w, h, fill=TINT)
            box = textbox(s, M + 0.2, y + 0.15, avail_w - 0.4, h - 0.3, lines, size=14)
            for p in box.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = "Menlo"
            y += h + 0.25
        elif kind == "quote":
            h = 1.5
            rect(s, M, y, avail_w, h, fill=TINT)
            rect(s, M, y, 0.12, h, fill=ACCENT)
            textbox(s, M + 0.4, y + 0.2, avail_w - 0.8, h - 0.4, [val], size=18, anchor="m")
            y += h + 0.25
    flush_bullets()
    textbox(s, W - M - 1.2, H - 0.55, 1.2, 0.35, [f"{page['n']} / {total}"], size=11, color=MUTED, align="c")
    if y > H - 0.4:
        print(f"  [溢出警告] 第 {page['n']} 頁「{page['title']}」內容估計高度到 {y:.1f} in")
    notes(s, page["notes"])
    return s


def build(key):
    md, out = DECKS[key]
    pages = parse(ROOT / md)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    label = {"09": "第四週 4-1", "10": "第四週 4-2", "11": "第四週 4-3"}[key]
    for p in pages:
        if p["n"] == 1:
            cover(prs, p, label)
        else:
            content(prs, p, len(pages))
    out_path = ROOT / "slides" / out
    prs.save(out_path)
    print(f"{out}: {len(pages)} 頁 ← {md}")


if __name__ == "__main__":
    keys = sys.argv[1:] or list(DECKS)
    for k in keys:
        build(k)
