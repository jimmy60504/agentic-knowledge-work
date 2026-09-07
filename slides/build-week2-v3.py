#!/usr/bin/env python3
"""第二週講師版投影片：依 drafts/02-week2-presentation-lesson-plan.md 建置。

用法：python3 slides/build-week2.py  → slides/week2-agent-knowledge-work-v3.pptx
版面依 drafts/07-visual-design-rules.md；配色另選（暖炭灰、陶土紅、鼠尾草綠）。
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image
import copy
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "assets"
OUT = ROOT / "slides" / "week2-agent-knowledge-work-v3.pptx"

# 配色：一色主導（白底＋炭灰字），一個強調色（陶土紅），一個輔色（鼠尾草綠）
DARK = RGBColor(0x2A, 0x26, 0x23)
ACCENT = RGBColor(0xB8, 0x50, 0x42)
SAGE = RGBColor(0xA7, 0xBE, 0xAE)
TINT = RGBColor(0xEE, 0xF2, 0xEF)
MUTED = RGBColor(0x6B, 0x65, 0x60)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD9, 0xD4, 0xCE)
FONT = "PingFang TC"

W, H = 13.333, 7.5
M = 0.6

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


# ---------- 基本元件 ----------
def _set_font(run, size, bold=False, color=DARK, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def tb(slide, x, y, w, h, lines, size=16, bold=False, color=DARK, align="l",
       anchor="t", spacing=6, margin=0):
    """lines: str 或 list；list 元素可為 str 或 dict(text,size,bold,color,bullet,italic)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(lines, str):
        lines = [lines]
    first = True
    for ln in lines:
        d = ln if isinstance(ln, dict) else {"text": ln}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[d.get("align", align)]
        p.space_after = Pt(d.get("after", spacing))
        if d.get("bullet"):
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(Inches(0.22))))
            pPr.set("indent", str(-int(Inches(0.22))))
            bu = etree.SubElement(pPr, qn("a:buChar"))
            bu.set("char", "•")
        r = p.add_run()
        r.text = d["text"]
        _set_font(r, d.get("size", size), d.get("bold", bold), d.get("color", color), d.get("italic", False))
    return box


def rect(slide, x, y, w, h, fill=TINT, line=None, rounded=True, radius=0.08):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    shp.shadow.inherit = False
    shp.text_frame.text = ""
    return shp


def circle_num(slide, x, y, n, d=0.5, fill=ACCENT, color=WHITE, size=16):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    _set_font(r, size, True, color)
    return c


def picture(slide, path, x, y, w, h, align="c"):
    """置入圖片並保持比例，置於 (x,y,w,h) 框內。"""
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / ih
    if w / h > ar:
        ph = h
        pw = h * ar
    else:
        pw = w
        ph = w / ar
    px = x + (w - pw) / 2 if align == "c" else x
    py = y + (h - ph) / 2
    pic = slide.shapes.add_picture(str(path), Inches(px), Inches(py), Inches(pw), Inches(ph))
    return pic, (px, py, pw, ph)


def arrow(slide, x1, y1, x2, y2, color=MUTED, dashed=False, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dashed:
        d = etree.SubElement(ln, qn("a:prstDash"))
        d.set("val", "dash")
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("h", "med")
    return c


def table(slide, x, y, w, rows, col_w, size=13, header=True, row_h=0.42):
    nrows, ncols = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    t = gt.table
    tblPr = t._tbl.tblPr
    tblPr.set("bandRow", "0")
    tblPr.set("firstRow", "0")
    # 移除預設樣式 id，改手動上色
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    for i, cw in enumerate(col_w):
        t.columns[i].width = Inches(cw)
    for ri, row in enumerate(rows):
        t.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            is_head = header and ri == 0
            cell.fill.fore_color.rgb = DARK if is_head else (WHITE if ri % 2 else TINT)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = str(val)
            _set_font(r, size, is_head or ci == 0, WHITE if is_head else DARK)
    return gt


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def background(slide, color):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = color


def content_slide(section, num, title, minutes=None):
    """白底內容頁：左上章節圓號＋段落標，標題 36pt。"""
    s = prs.slides.add_slide(BLANK)
    background(s, WHITE)
    if num is not None:
        circle_num(s, M, 0.55, num, d=0.46, size=15)
        tb(s, M + 0.6, 0.55, 6, 0.46, section, size=12, color=MUTED, anchor="m")
    else:
        tb(s, M, 0.55, 6, 0.46, section, size=12, color=MUTED, anchor="m")
    if minutes:
        tb(s, W - M - 2.2, 0.55, 2.2, 0.46, minutes, size=12, color=MUTED, align="r", anchor="m")
    tb(s, M, 1.15, W - 2 * M, 1.0, title, size=36, bold=True, anchor="t")
    return s


def dark_slide():
    s = prs.slides.add_slide(BLANK)
    background(s, DARK)
    return s


def card(slide, x, y, w, h, head, body, num=None, head_size=18, body_size=14):
    rect(slide, x, y, w, h)
    hx = x + 0.25
    if num is not None:
        circle_num(slide, x + 0.25, y + 0.25, num, d=0.42, size=14)
        hx = x + 0.8
    tb(slide, hx, y + 0.22, w - (hx - x) - 0.25, 0.5, head, size=head_size, bold=True, anchor="m")
    tb(slide, x + 0.25, y + 0.85, w - 0.5, h - 1.0, body, size=body_size, color=DARK, spacing=4)


# ============ 投影片 ============

# 1 標題
s = dark_slide()
tb(s, M, 1.4, 8, 0.5, "Agent 時代的知識工作｜第二週", size=16, color=SAGE)
tb(s, M, 2.1, 11, 2.2, ["先體驗 agent", "讓 agent 把你的想法做成成果"], size=44, bold=True, color=WHITE, spacing=10)
tb(s, M, 5.2, 11, 1.2, ["前半段 30 分：agent 能幫什麼、沒有判斷的產出長什麼樣",
                        "示範 30 分：講師走完整條線　　練習 120 分：自選題目，做出 10 到 12 頁"],
   size=16, color=SAGE, spacing=6)
circle_num(s, W - M - 0.8, 6.2, "2", d=0.8, fill=ACCENT, size=30)
notes(s, "本週核心：讓 agent 把你的想法做成成果。前半段講 agent 能協助什麼、如何合作，並用一份直接生成的簡報說明沒有判斷的產出長什麼樣；示範段由講師用預設題目走完整條線；練習段自選題目走七步。")

# 2 本週安排
s = content_slide("本週安排", None, "三段：講、看、做")
cols = [("前半段 30 分", "agent 能幫什麼", ["研究專案情境", "能力互補與四種協助", "請教與交辦", "對照組：漂亮但答不出為什麼", "兩個迴圈與三種產出", "保存過程"]),
        ("示範 30 分", "講師走完整條線", ["預設題：用 agent 做業務簡報，該選哪種方式", "七步走一遍，只展開一條論點", "先做一頁看畫面，其餘交給 agent", "最後展示完整十頁"]),
        ("練習 120 分", "自選題目，走七步", ["題目自選，預設題可用", "開頭先讓 agent 做第一版", "兩個檢查點與保底路徑", "成品 10 到 12 頁可編輯 PPTX"])]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, h2, items) in enumerate(cols):
    x = M + i * (cw + 0.3)
    rect(s, x, 2.4, cw, 4.3)
    tb(s, x + 0.3, 2.65, cw - 0.6, 0.5, h1, size=22, bold=True, color=ACCENT)
    tb(s, x + 0.3, 3.2, cw - 0.6, 0.5, h2, size=16, bold=True)
    tb(s, x + 0.3, 3.85, cw - 0.6, 2.7, [{"text": t, "bullet": True} for t in items], size=14, spacing=5)
notes(s, "不以視覺華麗程度評分，不要求 GitHub push 或課後作業。成品每一頁都要對應 draft 裡的一段內容，被問到答得出為什麼。")

# 3 開場：研究專案
s = content_slide("開場", None, "假設你正在推進一個研究專案", "3 分")
items = [("材料", "幾篇論文、一批資料、之前試過的方法"), ("想法", "一邊讀、一邊分析，一邊調整；有些問題還沒想清楚"), ("成果", "整理資料、做小工具、向同事說明目前的發現")]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, body) in enumerate(items):
    x = M + i * (cw + 0.3)
    card(s, x, 2.5, cw, 2.2, h1, body, num=i + 1, head_size=20, body_size=15)
tb(s, M, 5.1, W - 2 * M, 0.6, "Agent 可以參與這些工作：查找、整理、試作或檢查；採用什麼、發現什麼問題，一起保存，下次接著做。", size=16)
rect(s, M, 5.9, W - 2 * M, 0.9, fill=DARK)
tb(s, M + 0.3, 5.9, W - 2 * M - 0.6, 0.9, "要讓 agent 幫上忙，先得知道目前卡在哪裡、希望它協助哪一段。", size=20, bold=True, color=WHITE, anchor="m")
notes(s, "假設你正在做一個研究專案。手上有幾篇論文、一批資料和之前試過的方法，也有一些尚未想清楚的問題。你一邊讀材料、做分析，一邊調整想法；過程中還需要整理資料、製作小工具，或向同事說明目前的發現。Agent 可以參與這些工作。把目前的材料、想法與問題放進專案，讓它依需要協助查找、整理、試作或檢查；討論後採用哪些建議、發現什麼問題，也一起保存，下次從這裡接著做。收在：要讓 agent 幫上忙，先得知道目前卡在哪裡。")

# 4 AI 放大器
s = content_slide("一、AI 如何配合自己的能力與需要", 1, "領域知識 × 駕馭 agent 的能力", "7 分")
picture(s, ASSETS / "week2-evidence-pilot/magnifying-glass-book.jpg", M, 2.4, 4.6, 3.6, align="l")
tb(s, M, 6.1, 4.6, 0.5, "圖片：Julo，Wikimedia Commons，公有領域，未修改。", size=10, color=MUTED)
x = 5.7
card(s, x, 2.4, W - M - x, 1.9, "領域知識：理解方法與結果的因果關係", "知道一件事為什麼做得成，才看得出目前缺什麼、該調整哪個環節，再把材料、方法與預期結果向 agent 說清楚。", head_size=17, body_size=14)
card(s, x, 4.5, W - M - x, 1.9, "駕馭 agent 的能力：熟悉模型與工具", "知道不同模型擅長什麼、有哪些功能與限制，例如何時分工給 sub-agent，才能選擇合適的做法。", head_size=17, body_size=14)
tb(s, x, 6.55, W - M - x, 0.4, "乘號是兩種能力配合的比喻：一邊是零，成果也是零。", size=13, color=MUTED)
notes(s, "同樣在做研究，有人熟悉分析方法但不熟工具，有人能很快做出程式，卻不清楚結果該怎麼解讀。可以用「領域知識 × 駕馭 agent 的能力」理解 AI 放大器。例子：想做資料展示，知道要回答什麼問題、需要哪些資料、怎樣呈現才不會誤導，是領域判斷；知道可以請 agent 做互動網頁、如何試用與回報問題，是工具熟悉度。兩者配合，想法才能轉成具體可執行的要求。")

# 5 四種協助
s = content_slide("一、AI 如何配合自己的能力與需要", 1, "先從目前的需要選一個切入點")
helps = [("缺知識與方向", "請 agent 解釋、提供關鍵字與做法，再查閱或試做。"),
         ("有想法，但不熟悉工具操作", "描述想法、請 agent 實作，再試用與修改。例如互動網頁或資料展示。"),
         ("簡單、重複的雜工", "把規則與完成條件說清楚，交給 agent 處理後檢查。例如批次改檔名、整理欄位。"),
         ("需要協助檢查", "請它找遺漏與盲點，再依材料和工作目的決定是否修改。")]
cw = (W - 2 * M - 0.3) / 2
for i, (h1, body) in enumerate(helps):
    x = M + (i % 2) * (cw + 0.3)
    y = 2.4 + (i // 2) * 2.05
    card(s, x, y, cw, 1.85, h1, body, num=i + 1, head_size=18, body_size=14)
tb(s, M, 6.55, W - 2 * M, 0.5, "互動：說一件要向同事說明的事，指出目前卡在哪裡、想先試哪種協助。這件事就是練習的題目。", size=15, bold=True, color=ACCENT)
notes(s, "Agent 可以協助補足不熟悉的部分，也可以接手既有工作中的一些環節。同一件工作可能需要幾種協助。請幾位學員說一件要向同事說明的事、卡在哪裡、想先試哪種協助；講師當場記下來，練習時就有題。選定之後，下一步是把需要向 agent 說清楚。")

# 6 請教與交辦
s = content_slide("二、請教與交辦，保留自己的判斷", 2, "還沒方向就請教，有了方向就交辦", "5 分")
cw = (W - 2 * M - 0.3) / 2
rect(s, M, 2.4, cw, 3.3)
tb(s, M + 0.3, 2.6, cw - 0.6, 0.5, "請教：像和不同領域的導師討論", size=18, bold=True)
tb(s, M + 0.3, 3.15, cw - 0.6, 1.2, "帶著自己的專業、經驗與問題，說目前怎麼想、為什麼，再從它的觀點繼續追問。建議要經資料或實作確認。", size=14)
tb(s, M + 0.3, 4.45, cw - 0.6, 1.1, "「我要向同事比較幾種用 agent 做簡報的方式，比較條件該怎麼定？」", size=14, color=ACCENT, bold=True)
x2 = M + cw + 0.3
rect(s, x2, 2.4, cw, 3.3)
tb(s, x2 + 0.3, 2.6, cw - 0.6, 0.5, "交辦：把方向與重要限制講清楚", size=18, bold=True)
tb(s, x2 + 0.3, 3.15, cw - 0.6, 1.2, "說明給誰用、做成什麼、範圍在哪，以及長度、格式或必須保留的內容。區分必要條件與偏好，具體做法讓 agent 提出。", size=14)
tb(s, x2 + 0.3, 4.45, cw - 0.6, 1.1, "「把確認的內容做成十頁左右可編輯的 PowerPoint，16:9，先試做一頁。」", size=14, color=ACCENT, bold=True)
tb(s, M, 6.0, W - 2 * M, 0.9, "兩者可以來回切換：討論後有了方向就試做，成果出現新問題再回頭討論。採用什麼、如何修改，仍由自己決定。", size=15)
notes(s, "還不知道該用什麼方法，先請教；已經知道要做什麼，就把要求交代清楚，請它動手。兩句例句就是示範時會打的字。")

# 7 引文
s = dark_slide()
tb(s, M + 0.5, 2.0, W - 2 * M - 1, 2.0, "「他人幫助，但我主導」", size=44, bold=True, color=WHITE, align="l", anchor="m")
tb(s, M + 0.5, 4.2, W - 2 * M - 1, 0.6, "孫以瀚，2026-05-11 講座，第 61 頁標題。同頁指出人的貢獻包括形成問題、篩選與判斷、後續執行；屬講者個人觀點。", size=14, color=SAGE)
tb(s, M + 0.5, 5.0, W - 2 * M - 1, 0.5, "ethics.moe.edu.tw/files/resource/lecture/20260511/20260511_lecture_20260508.pdf#page=61", size=11, color=MUTED)
notes(s, "一頁引文帶過。人的貢獻：形成問題、篩選與判斷、後續執行。")

# 8 對照組
s = content_slide("三、對照組", 3, "一句題目貼給 agent，十分鐘，12 頁", "7 分")
pic, (px, py, pw, ph) = picture(s, ASSETS / "week2-baseline/baseline-comparison-stars.png", M, 2.3, 7.4, 4.3, align="l")
tb(s, M, py + ph + 0.1, 7.4, 0.4, "課程試做，2026-09-07 由 agent 生成，未經任何查證；四種方式與星等皆出自模型既有知識。", size=10, color=MUTED)
x = 8.4
tb(s, x, 2.3, W - M - x, 0.5, "只給它一句話：", size=14, color=MUTED)
tb(s, x, 2.7, W - M - x, 1.4, "「幫我做一份介紹用 agent 做業務簡報有哪些方式、怎麼選的投影片，給同事看的」", size=15, bold=True, color=ACCENT)
tb(s, x, 4.2, W - M - x, 2.4, [{"text": "得到 12 頁：議程、四種方式各一頁、星等比較表、四個判斷問題、情境對照、提醒與回顧。", "after": 8},
                              {"text": "第一眼會覺得做得不錯。", "bold": True}], size=14)
notes(s, "先看一份把題目直接貼給 agent 做出來的簡報。十分鐘後得到 12 頁。第一眼會覺得做得不錯。")

# 9 拿去報會卡住
s = content_slide("三、對照組", 3, "問題出在拿去報的時候")
probs = [("第四種「資料串接 agent」是什麼？", "講的人自己也不知道，一被追問便無法回答。"),
         ("星等沒有依據", "比較表的評分出自模型既有知識，沒有查證，沒有來源。"),
         ("情境不是我們的", "全是客戶提案與品牌規範，沒有送審、列印與承辦人接手。")]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, body) in enumerate(probs):
    x = M + i * (cw + 0.3)
    card(s, x, 2.4, cw, 2.4, h1, body, num=i + 1, head_size=16, body_size=14)
rect(s, M, 5.2, W - 2 * M, 1.4, fill=DARK)
tb(s, M + 0.3, 5.2, W - 2 * M - 0.6, 1.4, ["整份通順、漂亮，卻沒有一頁是自己的判斷。", "它漂亮的原因不是模型品味，而是照一份寫成文字的版面規則做，做完又看過畫面。"], size=17, bold=False, color=WHITE, anchor="m", spacing=6)
notes(s, "第四種「資料串接 agent」是什麼，講的人自己也不知道；星等沒有依據；情境全是客戶提案與品牌規範。整份通順、漂亮，卻沒有一頁是自己的判斷。它漂亮的原因是照一份寫成文字的版面規則做，做完又看過畫面。")

# 10 版面可外包、內容判斷不行
s = content_slide("三、對照組", 3, "版面規則可以交給 agent，內容判斷不行")
cw = (W - 2 * M - 0.3) / 2
rect(s, M, 2.4, cw, 3.6)
tb(s, M + 0.3, 2.6, cw - 0.6, 0.5, "對照組", size=20, bold=True, color=MUTED)
tb(s, M + 0.3, 3.2, cw - 0.6, 2.7, [{"text": "12 頁", "bullet": True}, {"text": "照版面規則排，看過畫面", "bullet": True}, {"text": "內容出自模型既有知識", "bullet": True}, {"text": "沒有來源，答不出為什麼", "bullet": True, "color": ACCENT, "bold": True}], size=16, spacing=8)
x2 = M + cw + 0.3
rect(s, x2, 2.4, cw, 3.6, fill=WHITE, line=ACCENT)
tb(s, x2 + 0.3, 2.6, cw - 0.6, 0.5, "今天要做的", size=20, bold=True, color=ACCENT)
tb(s, x2 + 0.3, 3.2, cw - 0.6, 2.7, [{"text": "10 到 12 頁，頁數相同", "bullet": True}, {"text": "同一套版面規則交給 agent", "bullet": True}, {"text": "內容來自自己查過的材料與 draft", "bullet": True}, {"text": "每一頁有依據，答得出為什麼", "bullet": True, "color": ACCENT, "bold": True}], size=16, spacing=8)
tb(s, M, 6.3, W - 2 * M, 0.5, "差別只在每一頁有沒有依據。版面的問題，示範第六步排版時另有辦法處理。", size=15)
notes(s, "今天要做的簡報頁數相同、版面也用同一套規則，差別只在每一頁有沒有依據、答不答得出為什麼。")

# 11 三種形式
s = content_slide("三、對照組", 3, "示範題：用 agent 做業務簡報，該選哪種方式")
forms = [("圖片式", "image2-voltgo-city-3x2.jpg", "VoltGo City 新款電動機車發表會", "每頁是一張圖片，任何裝置都能看；內容固定在畫面裡"),
         ("HTML", "html-store-manager-30-60-90-3x2.jpg", "新任店長內訓簡報", "在瀏覽器裡播放，可以有互動；改動存在哪裡要看作法"),
         ("PPTX", "pptx-flowpilot-2026-q2-clean-3x2.jpg", "FlowPilot 2026 Q2 營運回顧", "用 PowerPoint 開啟，逐頁逐字可改；交接與列印走既有流程")]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, fn, cap, desc) in enumerate(forms):
    x = M + i * (cw + 0.3)
    rect(s, x, 2.3, cw, 4.25)
    tb(s, x + 0.25, 2.45, cw - 0.5, 0.5, h1, size=20, bold=True, color=ACCENT, anchor="m")
    picture(s, ASSETS / "week2-nesa-examples" / fn, x + 0.25, 3.0, cw - 0.5, 1.6, align="l")
    tb(s, x + 0.25, 4.65, cw - 0.5, 0.35, cap, size=10, color=MUTED)
    tb(s, x + 0.25, 5.05, cw - 0.5, 1.4, desc, size=14)
tb(s, M, 6.7, W - 2 * M, 0.4, "乍看都像簡報，差別在做好之後能改什麼、改了存在哪。來源：slidefirm/NESA-SLIDE README 的 demo 縮圖，MIT 授權，未修改。", size=11, color=MUTED)
notes(s, "先認識圖片、HTML、PPTX 三大類。原站把 HTML 再分兩型，本課合併為一類。練習時選預設題的人，把 repo 連結交給 agent，每類實際打開一份；只看範例，不查它的安裝方式與製作工具。")

# 12 兩個迴圈（流程圖）
s = content_slide("四、建立論點與確認內容的兩個迴圈", 4, "兩個迴圈，往返之後才排版", "6 分")
def fbox(x, y, w, h, text, fill=WHITE, line=LINE, size=12, bold=False, color=DARK):
    r = rect(s, x, y, w, h, fill=fill, line=line, radius=0.12)
    tb(s, x + 0.08, y, w - 0.16, h, text, size=size, bold=bold, color=color, align="c", anchor="m")
    return r
# 左群組：建立論點
gx, gy, gw, gh = M, 2.35, 5.3, 3.75
rect(s, gx, gy, gw, gh, fill=TINT)
tb(s, gx + 0.2, gy + 0.1, 3, 0.35, "建立論點", size=13, bold=True, color=ACCENT)
bw, bh = 4.7, 0.72
bx = gx + 0.3
ya, yb, yc = gy + 0.5, gy + 1.5, gy + 2.5
fbox(bx, ya, bw, bh, "帶著問題探索：蒐集文獻、案例與圖片")
fbox(bx, yb, bw, bh, "人通讀資料，先寫大綱；比較依據，加入洞見與取捨", bold=True)
fbox(bx, yc, bw, bh, "Agent 整理故事線，人確認重點與順序")
arrow(s, bx + bw / 2, ya + bh, bx + bw / 2, yb)
arrow(s, bx + bw / 2, yb + bh, bx + bw / 2, yc)
arrow(s, gx + 0.16, yc + bh / 2, gx + 0.16, ya + bh / 2 + 0.05, dashed=True)  # 補探索、調整想法
tb(s, gx + 0.3, gy + 3.35, 4.5, 0.32, "虛線：補探索、調整想法", size=10, color=MUTED)
# 右群組：確認內容
hx, hy, hw, hh = 7.0, 2.35, 5.3, 3.75
rect(s, hx, hy, hw, hh, fill=TINT)
tb(s, hx + 0.2, hy + 0.1, 3, 0.35, "確認內容", size=13, bold=True, color=ACCENT)
cx = hx + 0.3
fbox(cx, ya, bw, bh, "Agent 在主 draft 展開文章，沿論點完整說明")
fbox(cx, yb, bw, bh, "人充實內容：改成自己的話，查核與補缺口", bold=True)
fbox(cx, yc, bw, bh, "整體通讀與收斂：Agent 建議刪併，人決定取捨")
arrow(s, cx + bw / 2, ya + bh, cx + bw / 2, yb)
arrow(s, cx + bw / 2, yb + bh, cx + bw / 2, yc)
arrow(s, hx + 0.16, yc + bh / 2, hx + 0.16, ya + bh / 2 + 0.05, dashed=True)
tb(s, hx + 0.3, hy + 3.35, 2.0, 0.32, "虛線：依取捨修改草稿", size=10, color=MUTED)
# 群組間箭頭 C → D
arrow(s, bx + bw, yc + bh / 2, cx, ya + bh / 2, color=ACCENT, width=2)
# 底部：分頁與排版
fbox(M, 6.3, 5.8, 0.6, "轉成投影片分頁：選用素材，安排口述與畫面", fill=DARK, line=DARK, color=WHITE, bold=True)
fbox(7.0, 6.3, 5.7, 0.6, "Agent 排版；人試講、檢查與微調", fill=DARK, line=DARK, color=WHITE, bold=True)
arrow(s, cx + bw / 2, yc + bh, 7.0 + 5.7 / 2, 6.3, color=ACCENT, width=2)
arrow(s, 5.8, 6.6, 7.0, 6.6, color=ACCENT, width=2)
notes(s, "「建立論點」處理想講什麼、如何組織；「確認內容」把草稿充實、收斂成自己能說明的內容。讀圖三件事：人先通讀、自己列綱；新材料可以改變方向，回到受影響的段落修正，不必整圈重做；接近交付時逐步收斂，內容盡量在 Markdown 修定後再排版。")

# 13 讀圖三件事＋三種產出
s = content_slide("四、建立論點與確認內容的兩個迴圈", 4, "讀圖抓三件事，做完留下三樣東西")
cw = (W - 2 * M - 0.3) / 2
rect(s, M, 2.4, cw, 4.2)
tb(s, M + 0.3, 2.6, cw - 0.6, 0.5, "讀圖三件事", size=18, bold=True)
tb(s, M + 0.3, 3.2, cw - 0.6, 3.2, [
    {"text": "人先通讀、自己列綱。", "bold": True, "after": 2}, {"text": "加入洞見與取捨，再讓 agent 沿方向整理；卡住可請它試串後逐段確認理由。", "after": 10},
    {"text": "新材料可以改變方向。", "bold": True, "after": 2}, {"text": "回到受影響的論點或段落修正，不必整圈重做。", "after": 10},
    {"text": "接近交付時逐步收斂。", "bold": True, "after": 2}, {"text": "先處理正確性與理解問題，內容在 Markdown 修定後再排版。"}], size=14)
x2 = M + cw + 0.3
rect(s, x2, 2.4, cw, 4.2, fill=WHITE, line=LINE)
tb(s, x2 + 0.3, 2.6, cw - 0.6, 0.5, "三樣東西", size=18, bold=True)
outs = [("kb/", "搜尋到的材料與來源"), ("drafts/", "採用的論點、文章與分頁草稿；棄案另放一份並記理由"), ("PPT", "對外表達的成品")]
for i, (h1, body) in enumerate(outs):
    y = 3.25 + i * 1.05
    circle_num(s, x2 + 0.3, y, i + 1, d=0.42, size=14)
    tb(s, x2 + 0.9, y - 0.05, cw - 1.2, 0.45, h1, size=16, bold=True, color=ACCENT, anchor="m")
    tb(s, x2 + 0.9, y + 0.4, cw - 1.2, 0.6, body, size=13)
tb(s, M, 6.75, W - 2 * M, 0.4, "三者都找得到，這份簡報才接得下去。", size=14, color=MUTED)
notes(s, "走完會留下三樣東西：kb 保存搜尋到的材料與來源；drafts 保存採用的論點、文章與分頁草稿，棄案另放一份並記理由；PPT 是對外表達的成品。")

# 14 保存過程
s = content_slide("五、保存過程，才能接著做", 5, "討論到值得留的，就請 agent 寫成檔案", "2 分")
steps = [("請 agent 整理成 Markdown", "資料、想法、建議、取捨及未定事項，連同來源與圖片"), ("自己打開確認", "意思有沒有改變、來源是否保留、未定事項是否標明"), ("下次指定讀取", "從這裡接著做，不必重講一遍")]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, body) in enumerate(steps):
    x = M + i * (cw + 0.3)
    card(s, x, 2.5, cw, 2.3, h1, body, num=i + 1, head_size=17, body_size=14)
    if i < 2:
        arrow(s, x + cw + 0.02, 3.65, x + cw + 0.28, 3.65, color=ACCENT, width=2)
tb(s, M, 5.3, W - 2 * M, 0.6, "資料增加本身不代表內容成熟，收斂依靠取捨。", size=16, bold=True)
tb(s, M, 5.9, W - 2 * M, 0.6, "接下來直接看整條線怎麼走。", size=16, color=MUTED)
notes(s, "討論到值得留下的內容，就請 agent 整理成 Markdown，自己打開確認；下次指定讀取，再接著做。保留資料、想法、建議、取捨及未定事項，也保存來源與圖片，讓結論能回到依據。")

# 15 示範：開始前
s = dark_slide()
tb(s, M, 0.8, 6, 0.5, "示範 30 分｜開始前", size=14, color=SAGE)
tb(s, M, 1.3, W - 2 * M, 1.0, "動手前先確認四件事", size=36, bold=True, color=WHITE)
pre = [("資料使用設定", "所用 agent 工具的對話是否用於訓練；關閉訓練不代表資料不上雲"), ("repo 公開與否", "練習專案若放上 GitHub，是公開還是私有；private repo 也在雲端"), ("PPTX 製作路徑", "這台機器用什麼做出可編輯的 PPTX，講師告知，不必自己摸索"), ("怎麼看畫面", "沒有預覽軟體時，請 agent 用 PowerPoint 或 Keynote 匯出 PDF 或圖片")]
cw = (W - 2 * M - 0.3) / 2
for i, (h1, body) in enumerate(pre):
    x = M + (i % 2) * (cw + 0.3)
    y = 2.6 + (i // 2) * 1.9
    rect(s, x, y, cw, 1.7, fill=RGBColor(0x3A, 0x35, 0x31))
    circle_num(s, x + 0.25, y + 0.25, i + 1, d=0.42, size=14)
    tb(s, x + 0.8, y + 0.22, cw - 1.05, 0.5, h1, size=18, bold=True, color=WHITE, anchor="m")
    tb(s, x + 0.25, y + 0.85, cw - 0.5, 0.8, body, size=13, color=SAGE)
tb(s, M, 6.55, W - 2 * M, 0.5, "使用公開、虛構或確認適合提供給工具的材料；舊簡報也要檢查備註與附件。", size=13, color=SAGE)
notes(s, "一頁投影片講完，不展開。數發部手冊提醒避免將個人訊息與機密資料提供給公開或 API 串接的生成式 AI 服務（第 81 頁，4.3.6）。拿不準的材料改用講師範例。")

# 16 示範腳本
s = content_slide("示範 30 分", None, "講師用預設題走完七步，只展開一條論點")
rows = [["分鐘", "步驟", "示範內容"],
        ["2–4", "一、初始化", "建 AGENTS.md 與 kb/第一次討論.md，打開看，請 agent 讀取"],
        ["4–8", "二、材料", "課前查好的範例與工具筆記，現場請 agent 讀，補一個工作上的條件"],
        ["8–12", "三、論點", "先說立場，agent 試串，改一處講錯的定位，棄一項比較並記理由"],
        ["12–16", "四、文章與查證", "展開一段，抓空泛結尾補條件；挑一句斷言查證，前後並列"],
        ["16–18", "五、分頁", "給頁面模板，agent 出分頁草稿；每頁對應 draft 一段，第一頁不能省"],
        ["18–25", "六、PPTX", "先做一頁，匯出看，太空，提具體修改；交視覺規則檔，其餘頁背景做"],
        ["25–30", "七、試講與回看", "講第一頁 60 秒，展示完整十頁；回看三問，補一條要求進 AGENTS.md"]]
table(s, M, 2.3, W - 2 * M, rows, [1.1, 2.4, W - 2 * M - 3.5], size=13, row_h=0.5)
tb(s, M, 6.55, W - 2 * M, 0.5, "學員不動手，記兩件事：自己工作上會加哪個條件；看到哪一步最想先試。", size=14, color=ACCENT, bold=True)
notes(s, "示範只展開「內容常改」這一條論點，三種方式的完整比較留給練習時選預設題的人自己做。材料筆記課前查好放在 kb 裡，現場只讀取與補充，並說明搜尋這一步練習時自己做。其餘頁面交給 agent 背景製作，做不完用預備版。")

# 17 練習規則
s = content_slide("練習 120 分", None, "題目自選，成品 10 到 12 頁")
cw = (W - 2 * M - 0.3) / 2
rect(s, M, 2.4, cw, 4.2)
tb(s, M + 0.3, 2.6, cw - 0.6, 0.5, "題目：預設題可用，自選題三個條件", size=17, bold=True)
tb(s, M + 0.3, 3.2, cw - 0.6, 2.2, [{"text": "有明確的受眾與用途：給誰看、看完要做什麼決定", "bullet": True}, {"text": "材料公開或可虛構，不碰署內真實資料", "bullet": True}, {"text": "兩小時內找得到依據", "bullet": True}], size=14, spacing=8)
tb(s, M + 0.3, 5.3, cw - 0.6, 1.2, "適合的類型：向同仁介紹一個新工具或新做法；比較某個業務流程的幾種改法；向不熟的同事說明一項研究或分析的進度。", size=13, color=MUTED)
x2 = M + cw + 0.3
rect(s, x2, 2.4, cw, 4.2, fill=WHITE, line=ACCENT)
tb(s, x2 + 0.3, 2.6, cw - 0.6, 0.5, "開始時先做兩件事", size=17, bold=True, color=ACCENT)
circle_num(s, x2 + 0.3, 3.25, 1, d=0.42, size=14)
tb(s, x2 + 0.9, 3.2, cw - 1.2, 1.3, [{"text": "把題目一句話貼給 agent 做第一版", "bold": True, "after": 2}, "讓它在背景跑，另開一個視窗繼續七步。這一版最後回看時要用。"], size=14)
circle_num(s, x2 + 0.3, 4.75, 2, d=0.42, size=14)
tb(s, x2 + 0.9, 4.7, cw - 1.2, 1.5, [{"text": "確認資料設定與 repo 公開與否", "bold": True, "after": 2}, "成品：10 到 12 頁可編輯 PPTX，每一頁對應 draft 裡的一段，被問到答得出為什麼。"], size=14)
notes(s, "預設題是示範題，材料連結與備份都現成。自選題要符合三個條件。第 10 分鐘講師掃一輪，寫不出受眾與用途的改用預設題。")

# 18 練習時間表
s = content_slide("練習 120 分", None, "七步時間表與兩個檢查點")
rows = [["分鐘", "步驟", "要做到"],
        ["0–5", "開始前", "確認資料設定；把題目貼給 agent 做第一版，另開視窗"],
        ["5–10", "一、初始化", "AGENTS.md、kb/第一次討論.md 寫好題目、受眾、用途，讓 agent 讀過"],
        ["10–30", "二、材料", "請 agent 找材料，每份實際打開，補自己的條件，整理進 kb；30 分到就停止搜尋"],
        ["30–45", "三、論點", "自己先說立場，agent 試串，至少改一處。檢查點一：kb 有題目、受眾、用途、三份材料，加一條故事線"],
        ["45–65", "四、文章與查證", "agent 展開文章，人補自己的例子，挑一句斷言請 agent 對照來源"],
        ["65–75", "五、分頁", "依模板出分頁草稿，每頁對應 draft 一段，口述兩三句"],
        ["75–100", "六、PPTX", "先做一頁、看畫面、提修改、交規則、展開其餘頁。檢查點二：第一頁做好且看過畫面"],
        ["100–110", "七、試講", "兩人一組，各講最沒把握的一頁，聽的人重述"],
        ["110–120", "回看", "第一版與自己版並排；回答三問；補一條要求進 AGENTS.md"]]
table(s, M, 2.25, W - 2 * M, rows, [1.1, 1.9, W - 2 * M - 3.0], size=12, row_h=0.44)
tb(s, M, 6.75, W - 2 * M, 0.4, "檢查點沒過的人走保底路徑：交第 2 頁、主體任一頁、建議頁，加試講一頁。", size=13, color=ACCENT, bold=True)
notes(s, "兩個檢查點的用意是保底：45 分沒有故事線、95 分沒有一頁 PPTX，就啟用保底路徑，確保每個人都走完一條論點到試講。講師最需要介入的三節：第二節搜尋過頭或走到安裝與 skill；第四節第 50 分鐘全場再示範一次查證；第六節首次輸出太空時引導提具體修改，不重做。")

# 19 頁面模板
s = content_slide("練習 120 分", None, "頁面模板：每一頁都要對應 draft 裡的一段")
rows = [["頁", "內容"], ["1–2", "受眾需要什麼，這次用什麼條件或角度看"], ["3–7", "主體：各選項、各發現或各步驟，一項一頁，附真實範例"], ["8–9", "套入受眾的情境，說明取捨"], ["10", "建議與理由"], ["11", "這份簡報自己是怎麼做的"], ["12", "來源與備註，可併入各頁備註"]]
table(s, M, 2.3, 7.4, rows, [1.0, 6.4], size=14, row_h=0.5)
x = 8.5
card(s, x, 2.3, W - M - x, 1.7, "對不上的頁刪掉", "agent 自行補的議程頁、回顧頁、提醒頁多半屬於這類。", head_size=16, body_size=13)
card(s, x, 4.2, W - M - x, 1.7, "第一頁不能省", "留一小段開場，聽眾才知道在比什麼標準。", head_size=16, body_size=13)
notes(s, "從文章整理 10 到 12 頁的分頁草稿，逐頁決定重點、口述、畫面與節奏。口述寫兩三句重點與過渡，不是逐字稿。")

# 20 回看（結尾）
s = dark_slide()
tb(s, M, 0.8, 6, 0.5, "回看", size=14, color=SAGE)
tb(s, M, 1.3, W - 2 * M, 1.0, "哪一頁答得出為什麼？", size=40, bold=True, color=WHITE)
tb(s, M, 2.4, W - 2 * M, 0.6, "把開頭那份直接生成的第一版和自己做的並排，再回答三問：", size=16, color=SAGE)
qs = [("補上了什麼內容", "查到的材料、自己的例子、修正過的斷言"), ("哪一項是自己的判斷", "採用或刪除的論點、棄案的理由、開頭在意的條件"), ("哪些交給了 agent", "搜尋、整理、查證、排版")]
cw = (W - 2 * M - 0.6) / 3
for i, (h1, body) in enumerate(qs):
    x = M + i * (cw + 0.3)
    rect(s, x, 3.2, cw, 2.2, fill=RGBColor(0x3A, 0x35, 0x31))
    circle_num(s, x + 0.25, 3.45, i + 1, d=0.42, size=14)
    tb(s, x + 0.8, 3.42, cw - 1.05, 0.5, h1, size=18, bold=True, color=WHITE, anchor="m")
    tb(s, x + 0.25, 4.05, cw - 0.5, 1.2, body, size=14, color=SAGE)
tb(s, M, 5.8, W - 2 * M, 1.0, ["將一項值得沿用的要求補進 AGENTS.md，例如「查到範例做法時，先確認在這台機器上跑得通」。", "留下可回查的材料、自己的 draft 與可編輯投影片；下次想處理的工作，現在告訴講師。"], size=15, color=WHITE, spacing=6)
notes(s, "最後回到本週核心：讓 agent 把你的想法做成成果。這次留下可回查的材料、自己的 draft 與可編輯投影片，這些內容會繼續留在專案裡。講師收集下次想處理的工作。")

prs.save(OUT)
print(f"saved {OUT} ({len(prs.slides)} slides)")
