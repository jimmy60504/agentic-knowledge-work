#!/usr/bin/env python3
"""第七週投影片建置：把 drafts/22-week7-slides.md 產成 PPTX（slides/week7.pptx）。

用法：python3 slides/build-week7.py
草稿截圖：設定環境變數 WEEK7_DRAFT_IMAGES 指向講師本機的對照表（TSV，不進 repo），
格式同第六週：「頁碼<TAB>圖片路徑<TAB>圖說」，有對照的頁以截圖取代預留位置並標【草稿截圖】。

版面、配色與解析全部沿用 build-week6.py（再往上沿用 build-week5.py），這裡只換來源與輸出。
"""
import importlib.util
import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
spec = importlib.util.spec_from_file_location("bw6", HERE / "build-week6.py")
bw6 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bw6)

bw5 = bw6.bw5
bw5.SRC = bw5.ROOT / "drafts/22-week7-slides.md"
bw5.OUT = HERE / "week7.pptx"

if os.environ.get("WEEK7_DRAFT_IMAGES"):
    os.environ["WEEK6_DRAFT_IMAGES"] = os.environ["WEEK7_DRAFT_IMAGES"]
    bw6.DRAFT.clear()
    bw6.DRAFT.update(bw6.draft_images())

if os.environ.get("WEEK7_OUT"):
    bw5.OUT = Path(os.environ["WEEK7_OUT"])

# 「- 樣式：補充」：延伸介紹的頁面，換成淡藍灰底並在右上角標「補充資料」，與主線的頁面區隔
STYLE = "樣式："
SUPP_BG, SUPP_TAG = "EEF2F6", "5B7083"
_content = bw6.content


def content(prs, page, total, section=""):
    style = next((v[len(STYLE):].strip() for k, v in page["blocks"]
                  if k == "bullet" and v.startswith(STYLE)), "")
    page["blocks"] = [(k, v) for k, v in page["blocks"] if not (k == "bullet" and v.startswith(STYLE))]
    _content(prs, page, total, section)
    if style != "補充":
        return
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    s = prs.slides[-1]
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor.from_string(SUPP_BG)
    bg.line.fill.background()
    tree = s.shapes._spTree
    tree.remove(bg._element)
    tree.insert(2, bg._element)  # 移到最底層
    w, h = Inches(1.3), Inches(0.38)
    tag = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, prs.slide_width - Inches(0.6) - w, Inches(0.42), w, h)
    tag.fill.solid()
    tag.fill.fore_color.rgb = RGBColor.from_string(SUPP_TAG)
    tag.line.fill.background()
    tf = tag.text_frame
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = 2  # 置中
    r = p.add_run()
    r.text = "補充資料"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


bw5.content = content

if __name__ == "__main__":
    bw5.main()
